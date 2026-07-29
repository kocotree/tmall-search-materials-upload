from collections import Counter
import csv
from dataclasses import dataclass
import hashlib
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .models import ProductRecord


PRODUCT_REQUIRED_COLUMNS = {
    "商品ID",
    "商品名称（查找引用）",
    "货号（查找引用）",
    "产品等级",
    "链接",
    "运营",
    "组别",
    "品类-公司维度划分",
}

HEADER_ALIASES = {
    "商品id": "商品ID",
    "短标题": "商家短标题",
}


class SchemaError(ValueError):
    """Raised when an input table cannot be interpreted safely."""


@dataclass(frozen=True)
class ProductValidationReport:
    row_count: int
    reason_codes_by_row: dict[int, list[str]]
    batch_blocking: bool = False

    @property
    def blocking(self) -> bool:
        return self.batch_blocking or bool(self.reason_codes_by_row)

    @property
    def blocked_row_count(self) -> int:
        return len(self.reason_codes_by_row)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _canonical_header(value: object) -> str:
    header = _cell_text(value)
    return HEADER_ALIASES.get(header.casefold(), header)


def _assert_headers(headers: list[str], required: set[str]) -> None:
    populated = [header for header in headers if header]
    duplicates = sorted(header for header, count in Counter(populated).items() if count > 1)
    missing = sorted(required - set(populated))
    problems = []
    if missing:
        problems.append("缺少必需表头: " + ", ".join(missing))
    if duplicates:
        problems.append("存在重复表头: " + ", ".join(duplicates))
    if problems:
        raise SchemaError("; ".join(problems))


def _read_xlsx(path: Path, required: set[str]) -> list[dict[str, str]]:
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            header_row = next(iterator)
        except StopIteration as error:
            raise SchemaError("工作簿没有表头") from error
        headers = [_canonical_header(value) for value in header_row]
        _assert_headers(headers, required)
        rows = []
        for values in iterator:
            if not any(value not in (None, "") for value in values):
                continue
            rows.append(
                {
                    header: _cell_text(values[index] if index < len(values) else None)
                    for index, header in enumerate(headers)
                    if header
                }
            )
        return rows
    finally:
        workbook.close()


def read_basic_materials_xlsx(path: Path) -> list[dict[str, str]]:
    return _read_xlsx(Path(path), {"商品ID", "商品标题", "商品白底图", "商家短标题"})


def read_search_materials_xlsx(path: Path) -> list[dict[str, str]]:
    return _read_xlsx(Path(path), {"商品ID", "素材类型"})


def read_product_csv(path: Path) -> list[ProductRecord]:
    with Path(path).open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        try:
            raw_headers = next(reader)
        except StopIteration as error:
            raise SchemaError("CSV 没有表头") from error
        headers = [_canonical_header(value) for value in raw_headers]
        _assert_headers(headers, PRODUCT_REQUIRED_COLUMNS)
        records = []
        for row_number, values in enumerate(reader, 2):
            if not any(value.strip() for value in values):
                continue
            raw = {
                header: values[index].strip() if index < len(values) else ""
                for index, header in enumerate(headers)
                if header
            }
            raw["_source_row"] = row_number
            records.append(
                ProductRecord(
                    product_id=raw["商品ID"],
                    sku=raw["货号（查找引用）"],
                    title=raw["商品名称（查找引用）"],
                    grade=raw["产品等级"],
                    link=raw["链接"],
                    owner=raw["运营"],
                    team=raw["组别"],
                    category=raw["品类-公司维度划分"],
                    raw=raw,
                )
            )
        return records


def validate_product_records(records: Iterable[ProductRecord]) -> ProductValidationReport:
    rows = list(records)
    id_counts = Counter(record.product_id for record in rows if record.product_id)
    reasons_by_row: dict[int, list[str]] = {}
    for fallback_row, record in enumerate(rows, 2):
        row_number = int(record.raw.get("_source_row", fallback_row))
        reasons = []
        if not record.product_id:
            reasons.append("MISSING_PRODUCT_ID")
        elif not record.product_id.isdigit():
            reasons.append("INVALID_PRODUCT_ID")
        elif id_counts[record.product_id] > 1:
            reasons.append("DUPLICATE_PRODUCT_ID")
        if reasons:
            reasons_by_row[row_number] = reasons
    return ProductValidationReport(
        row_count=len(rows),
        reason_codes_by_row=reasons_by_row,
    )
