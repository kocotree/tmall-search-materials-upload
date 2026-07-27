import argparse
import csv
import json
import sqlite3
import stat
from contextlib import contextmanager
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook
from PIL import Image

import pytest
import upload_search_materials.cli as cli_module
from upload_search_materials.approval import create_manifest

from upload_search_materials.cli import (
    BrowserSessionRequired,
    build_parser,
    main,
    page_context,
    partition_persisted_items,
)
from upload_search_materials.interaction.session import SessionStore as InteractionSessionStore
from upload_search_materials.interaction.stages import STAGES
from upload_search_materials.reporting import material_item_from_dict
from upload_search_materials.state_store import StateStore
from upload_search_materials.browser.upload_page import UploadOutcome


PRODUCT_HEADERS = [
    "商品ID",
    "商品名称（查找引用）",
    "货号（查找引用）",
    "产品等级",
    "链接",
    "运营",
    "组别",
    "品类-公司维度划分",
]


class BrowserSentinel:
    def __init__(self):
        self.publish_page_open_count = 0


class InteractionAppSentinel:
    def __init__(self):
        self.run_calls = []

    def run(self, **kwargs):
        self.run_calls.append(kwargs)


def test_interaction_commands_are_exposed_with_registry_stage_choices():
    parser = build_parser()

    interact = parser.parse_args(["interact", "--runs-root", "runs"])
    wait = parser.parse_args(
        [
            "wait-handoff",
            "--runs-root", "runs",
            "--session", "20260721_143025",
            "--stage", "setup",
        ]
    )
    wait_parser = next(
        action.choices["wait-handoff"]
        for action in parser._actions
        if isinstance(action, argparse._SubParsersAction)
    )
    stage_action = next(action for action in wait_parser._actions if action.dest == "stage")

    assert interact.command == "interact"
    assert interact.port == 8765
    assert wait.command == "wait-handoff"
    assert wait.timeout is None
    assert tuple(stage_action.choices) == tuple(stage.id for stage in STAGES)


def test_inspect_completeness_cli_writes_stage_two_contract(tmp_path):
    products = tmp_path / "products.csv"
    with products.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=PRODUCT_HEADERS)
        writer.writeheader()
        writer.writerow(
            {
                "商品ID": "1",
                "商品名称（查找引用）": "测试商品",
                "货号（查找引用）": "KQ001",
                "产品等级": "A",
                "链接": "https://example.invalid/1",
                "运营": "tester",
                "组别": "test",
                "品类-公司维度划分": "测试",
            }
        )
    promotion = tmp_path / "promotion.csv"
    with promotion.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=["商品ID", "目标容量", "现有素材数", "缺失数量", "状态", "证据"],
        )
        writer.writeheader()
        writer.writerow(
            {
                "商品ID": "1",
                "目标容量": "3",
                "现有素材数": "1",
                "缺失数量": "2",
                "状态": "needs_manual_review",
                "证据": "source=test",
            }
        )
    output = tmp_path / "completeness.json"

    code = main(
        [
            "inspect-completeness",
            "--products", str(products),
            "--promotion-status", str(promotion),
            "--output", str(output),
        ]
    )

    payload = json.loads(output.read_text(encoding="utf-8"))
    assert code == 0
    assert payload["contract_version"] == 1
    assert payload["products"][0]["promotion"]["missing_count"] == 2


def test_interact_creates_one_session_and_serves_its_url(tmp_path, monkeypatch, capsys):
    app = InteractionAppSentinel()
    calls = {"create": 0, "load": []}

    class FakeStore:
        def __init__(self, runs_root):
            assert Path(runs_root) == tmp_path

        def create_session(self):
            calls["create"] += 1
            return type("Session", (), {"session_id": "20260722_101112"})()

        def load_session(self, session_id):
            calls["load"].append(session_id)

    monkeypatch.setattr(cli_module, "SessionStore", FakeStore)
    monkeypatch.setattr(cli_module, "create_app", lambda runs_root, runtime_config=None: app)

    code = main(["interact", "--runs-root", str(tmp_path), "--port", "9123"])

    assert code == 0
    assert calls == {"create": 1, "load": []}
    assert app.run_calls == [
        {"host": "127.0.0.1", "port": 9123, "debug": False, "use_reloader": False}
    ]
    assert "http://127.0.0.1:9123/?session_id=20260722_101112" in capsys.readouterr().out


def test_interact_resumes_explicit_session_without_creating_another(
    tmp_path, monkeypatch, capsys
):
    app = InteractionAppSentinel()
    calls = {"create": 0, "load": []}

    class FakeStore:
        def __init__(self, runs_root):
            assert Path(runs_root) == tmp_path

        def create_session(self):
            calls["create"] += 1
            raise AssertionError("explicit session must not create a session")

        def load_session(self, session_id):
            calls["load"].append(session_id)
            return {"session_id": session_id}

    monkeypatch.setattr(cli_module, "SessionStore", FakeStore)
    monkeypatch.setattr(cli_module, "create_app", lambda runs_root, runtime_config=None: app)

    code = main(
        [
            "interact",
            "--runs-root", str(tmp_path),
            "--session", "20260722_101112",
        ]
    )

    assert code == 0
    assert calls == {"create": 0, "load": ["20260722_101112"]}
    assert "session_id=20260722_101112" in capsys.readouterr().out


def test_interact_uses_environment_only_when_runs_root_is_absent(
    tmp_path, monkeypatch
):
    environment_root = tmp_path / "environment"
    explicit_root = tmp_path / "explicit"
    roots = []
    app = InteractionAppSentinel()

    class FakeStore:
        def __init__(self, runs_root):
            roots.append(Path(runs_root))

        def create_session(self):
            return type("Session", (), {"session_id": "20260722_101112"})()

    monkeypatch.setenv("TMALL_RUNS_ROOT", str(environment_root))
    monkeypatch.setattr(cli_module, "SessionStore", FakeStore)
    monkeypatch.setattr(cli_module, "create_app", lambda runs_root, runtime_config=None: app)

    assert main(["interact"]) == 0
    assert main(["interact", "--runs-root", str(explicit_root)]) == 0

    assert roots == [environment_root, explicit_root]


def test_interact_rejects_occupied_port_before_creating_session(
    tmp_path, monkeypatch, capsys
):
    created = []

    class StoreMustNotStart:
        def __init__(self, runs_root):
            created.append(Path(runs_root))

    monkeypatch.setattr(cli_module, "_port_is_available", lambda port: False)
    monkeypatch.setattr(cli_module, "_port_owner_pid", lambda port: 4321)
    monkeypatch.setattr(cli_module, "SessionStore", StoreMustNotStart)

    assert main(["interact", "--runs-root", str(tmp_path), "--port", "8768"]) == 2
    assert created == []
    error = capsys.readouterr().err
    assert "8768" in error and "PID 4321" in error


def test_interact_defaults_to_runtime_project_runs_root(tmp_path, monkeypatch):
    roots = []
    app = InteractionAppSentinel()
    runtime = SimpleNamespace(runs_root=tmp_path / "runs")

    class FakeStore:
        def __init__(self, runs_root):
            roots.append(Path(runs_root))

        def create_session(self):
            return type("Session", (), {"session_id": "20260724_180000"})()

    monkeypatch.delenv("TMALL_RUNS_ROOT", raising=False)
    monkeypatch.setattr(cli_module, "load_runtime_config", lambda config=None: runtime)
    monkeypatch.setattr(cli_module, "SessionStore", FakeStore)
    monkeypatch.setattr(
        cli_module, "create_app", lambda runs_root, runtime_config=None: app
    )

    assert main(["interact"]) == 0
    assert roots == [tmp_path / "runs"]


def test_wait_handoff_prints_validated_submission(tmp_path, capsys):
    store = InteractionSessionStore(tmp_path)
    session = store.create_session()
    expected = store.save_input(session.session_id, "setup", {"store": "测试店铺"})

    code = main(
        [
            "wait-handoff",
            "--runs-root", str(tmp_path),
            "--session", session.session_id,
            "--stage", "setup",
            "--timeout", "1",
        ]
    )

    assert code == 0
    assert json.loads(capsys.readouterr().out) == expected


def test_wait_handoff_timeout_returns_two_without_business_execution(
    tmp_path, monkeypatch, capsys
):
    store = InteractionSessionStore(tmp_path)
    session = store.create_session()
    monkeypatch.setattr(
        cli_module,
        "_run",
        lambda args: (_ for _ in ()).throw(AssertionError("business execution invoked")),
    )
    monkeypatch.setattr(
        cli_module,
        "_publish",
        lambda *args, **kwargs: (_ for _ in ()).throw(
            AssertionError("publish execution invoked")
        ),
    )

    code = main(
        [
            "wait-handoff",
            "--runs-root", str(tmp_path),
            "--session", session.session_id,
            "--stage", "setup",
            "--timeout", "0",
        ]
    )

    captured = capsys.readouterr()
    assert code == 2
    assert captured.out == ""
    assert "timed out waiting for handoff" in captured.err


class CliFakeLocator:
    def __init__(self, page, selector):
        self.page = page
        self.selector = selector

    def count(self):
        if self.selector in self.page.counts:
            return self.page.counts[self.selector]
        return 1 if self.selector in self.page.texts else 0

    def inner_text(self):
        return self.page.texts.get(self.selector, "")

    def all_inner_texts(self):
        return list(self.page.lists.get(self.selector, []))

    def is_visible(self):
        return self.selector in self.page.visible

    def fill(self, value):
        self.page.filled.append((self.selector, value))

    def press(self, key):
        self.page.pressed.append((self.selector, key))

    def click(self):
        self.page.clicked.append(self.selector)


class CliFakeDownload:
    def __init__(self, source, suggested_filename):
        self.source = Path(source)
        self.suggested_filename = suggested_filename

    def save_as(self, destination):
        Path(destination).write_bytes(self.source.read_bytes())


class CliFakeDownloadInfo:
    def __init__(self, value):
        self.value = value


class CliFakePage:
    def __init__(self, downloads=None):
        self.downloads = iter(downloads or [])
        self.texts = {"#store": "KK Tree"}
        self.lists = {}
        self.counts = {}
        self.visible = set()
        self.filled = []
        self.pressed = []
        self.clicked = []

    def locator(self, selector):
        return CliFakeLocator(self, selector)

    @contextmanager
    def expect_download(self):
        yield CliFakeDownloadInfo(next(self.downloads))


def write_sources(tmp_path):
    products = tmp_path / "products.csv"
    with products.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=PRODUCT_HEADERS)
        writer.writeheader()
        writer.writerow(
            {
                "商品ID": "123",
                "商品名称（查找引用）": "普通水杯",
                "货号（查找引用）": "SKU-1",
                "产品等级": "A级",
                "链接": "https://example.invalid/123",
                "运营": "小王",
                "组别": "一组",
                "品类-公司维度划分": "水杯",
            }
        )
    rules = tmp_path / "rules.csv"
    with rules.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["月份", "品类", "要推等级"])
        writer.writeheader()
        writer.writerow({"月份": "7月", "品类": "水杯", "要推等级": "A级"})
    basic = tmp_path / "basic.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["商品ID", "商品标题", "商品白底图", "短标题"])
    sheet.append([123, "普通水杯", "已完成", "水杯"])
    workbook.save(basic)
    search = tmp_path / "search.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["商品id", "商品名称", "素材id", "素材类型"])
    sheet.append([123, "普通水杯", "", ""])
    workbook.save(search)
    return products, rules, basic, search


def test_dry_run_never_opens_publish_page(tmp_path):
    products, rules, basic, search = write_sources(tmp_path)
    output = tmp_path / "run"
    browser = BrowserSentinel()

    exit_code = main(
        [
            "run",
            "--mode", "dry-run",
            "--month", "7",
            "--store", "KK Tree",
            "--products", str(products),
            "--rules", str(rules),
            "--basic", str(basic),
            "--search", str(search),
            "--output", str(output),
            "--started-at", "2026-07-17T10:00:00+08:00",
        ],
        page=browser,
    )

    assert exit_code == 0
    assert browser.publish_page_open_count == 0
    assert (output / "review.html").exists()
    assert (output / "eligibility.csv").exists()
    assert (output / "product-tasks.json").exists()
    assert (output / "material-items.json").exists()
    assert (output / "supplement-candidates.csv").exists()
    assert (output / "run.sqlite3").exists()


def test_run_defaults_to_dry_run(tmp_path):
    products, rules, basic, search = write_sources(tmp_path)
    output = tmp_path / "run"

    exit_code = main(
        [
            "run",
            "--month", "7",
            "--store", "KK Tree",
            "--products", str(products),
            "--rules", str(rules),
            "--basic", str(basic),
            "--search", str(search),
            "--output", str(output),
            "--started-at", "2026-07-17T10:00:00+08:00",
        ]
    )

    metadata = json.loads((output / "run.json").read_text(encoding="utf-8"))
    assert exit_code == 0
    assert metadata["mode"] == "dry-run"


def test_approve_creates_manifest_only_for_selected_ready_items(tmp_path):
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    (run_dir / "run.json").write_text(
        json.dumps(
            {
                "store": "KK Tree",
                "run_id": "RUN-1",
                "source_sha256": {"products": "a" * 64},
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    (run_dir / "product-tasks.json").write_text(
        json.dumps([{"product_id": "123", "status": "ready_for_review"}]),
        encoding="utf-8",
    )
    (run_dir / "material-items.json").write_text(
        json.dumps(
            [
                {
                    "task_id": "MAT-1",
                    "product_id": "123",
                    "material_type": "image_text",
                    "slot_index": 1,
                    "status": "ready_for_review",
                    "assets": [
                        {
                            "product_id": "123",
                            "source_path": "a.png",
                            "sha256": "a" * 64,
                            "license_status": "confirmed",
                            "validation_status": "valid",
                        }
                    ],
                    "title": "KK树便携水杯",
                    "description": "便携水杯设计，满足日常携带和饮水使用需求。",
                    "content_hash": "hash-1",
                    "reason_codes": [],
                }
            ],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    exit_code = main(
        [
            "approve",
            "--run-dir", str(run_dir),
            "--task-id", "MAT-1",
            "--confirmed-by", "operator",
            "--confirmed-at", "2026-07-17T10:00:00+08:00",
            "--valid-until", "2026-07-18T10:00:00+08:00",
        ]
    )
    manifest = json.loads((run_dir / "approval-manifest.json").read_text(encoding="utf-8"))
    persisted_items = json.loads((run_dir / "material-items.json").read_text(encoding="utf-8"))
    state = StateStore(run_dir / "run.sqlite3")

    assert exit_code == 0
    assert manifest["entries"][0]["task_id"] == "MAT-1"
    assert manifest["run_id"] == "RUN-1"
    assert manifest["source_sha256"] == {"products": "a" * 64}
    assert len(manifest["manifest_sha256"]) == 64
    assert persisted_items[0]["status"] == "approved"
    assert state.item_status("MAT-1") == "approved"
    assert state.transitions_for("MAT-1")[0]["reason"] == "APPROVED_BY_USER"
    state.close()


def test_approve_rejects_child_of_blocked_product(tmp_path):
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    (run_dir / "run.json").write_text(
        json.dumps({"store": "KK Tree", "run_id": "RUN-1", "source_sha256": {}}),
        encoding="utf-8",
    )
    (run_dir / "product-tasks.json").write_text(
        json.dumps([{"product_id": "123", "status": "blocked"}]),
        encoding="utf-8",
    )
    (run_dir / "material-items.json").write_text(
        json.dumps(
            [
                {
                    "task_id": "MAT-1",
                    "product_id": "123",
                    "material_type": "image_text",
                    "slot_index": 1,
                    "status": "ready_for_review",
                    "assets": [],
                    "title": "标题",
                    "description": "描述内容满足长度要求。",
                }
            ]
        ),
        encoding="utf-8",
    )

    exit_code = main(
        [
            "approve",
            "--run-dir", str(run_dir),
            "--task-id", "MAT-1",
            "--confirmed-by", "operator",
            "--confirmed-at", "2026-07-17T10:00:00+08:00",
            "--valid-until", "2026-07-18T10:00:00+08:00",
        ]
    )

    assert exit_code == 1
    assert not (run_dir / "approval-manifest.json").exists()


def test_publish_requires_manifest_before_browser_use(tmp_path):
    browser = BrowserSentinel()

    exit_code = main(
        [
            "publish",
            "--run-dir", str(tmp_path),
            "--store", "KK Tree",
            "--selectors", str(tmp_path / "selectors.yaml"),
        ],
        page=browser,
    )

    assert exit_code == 2
    assert browser.publish_page_open_count == 0


def test_report_summarizes_existing_artifacts(tmp_path):
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    (run_dir / "run.json").write_text(
        json.dumps(
            {
                "run_id": "RUN-1",
                "store": "KK Tree",
                "month": 7,
                "mode": "dry-run",
                "source_sha256": {"products": "a" * 64},
            }
        ),
        encoding="utf-8",
    )
    (run_dir / "product-tasks.json").write_text(
        json.dumps([{"status": "blocked"}, {"status": "ready_for_review"}]),
        encoding="utf-8",
    )
    (run_dir / "material-items.json").write_text(
        json.dumps([{"task_id": "MAT-1", "status": "approved"}]),
        encoding="utf-8",
    )
    (run_dir / "upload-results.json").write_text(
        json.dumps(
            [
                {
                    "task_id": "MAT-1",
                    "status": "under_review",
                    "remote_material_id": "RM-1",
                    "reason": "",
                }
            ]
        ),
        encoding="utf-8",
    )

    exit_code = main(["report", "--run-dir", str(run_dir)])
    summary = (run_dir / "summary.md").read_text(encoding="utf-8")

    assert exit_code == 0
    assert "blocked: 1" in summary
    assert "ready_for_review: 1" in summary
    assert "目标月份: 7" in summary
    assert "products: " + "a" * 64 in summary
    assert "under_review: 1" in summary
    assert "MAT-1 / RM-1 / under_review" in summary


def test_run_builds_reviewable_items_when_all_inputs_are_resolved(tmp_path):
    products, rules, basic, search = write_sources(tmp_path)
    backend = tmp_path / "backend.csv"
    with backend.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=["商品ID", "目标坑位", "空坑位", "审核状态完整", "采集时间", "证据"],
        )
        writer.writeheader()
        writer.writerow(
            {
                "商品ID": "123",
                "目标坑位": "3",
                "空坑位": "1;2;3",
                "审核状态完整": "true",
                "采集时间": "2026-07-17T09:00:00+08:00",
                "证据": "fixture",
            }
        )
    asset_dir = tmp_path / "assets" / "123"
    asset_dir.mkdir(parents=True)
    for index in range(9):
        Image.new("RGB", (400, 400), color=(index, index, index)).save(asset_dir / f"{index}.png")
    copy_responses = tmp_path / "copy-responses.json"
    copy_responses.write_text(
        json.dumps(
            {
                f"123:{slot}": {
                    "title": f"KK树便携水杯{slot}",
                    "description": "便携水杯设计，满足日常携带和饮水使用需求。",
                }
                for slot in (1, 2, 3)
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    output = tmp_path / "resolved-run"

    exit_code = main(
        [
            "run",
            "--month", "7",
            "--store", "KK Tree",
            "--products", str(products),
            "--rules", str(rules),
            "--basic", str(basic),
            "--search", str(search),
            "--backend-status", str(backend),
            "--asset-root", str(tmp_path / "assets"),
            "--license-status", "confirmed",
            "--copy-responses", str(copy_responses),
            "--output", str(output),
            "--started-at", "2026-07-17T10:00:00+08:00",
        ]
    )
    items = json.loads((output / "material-items.json").read_text(encoding="utf-8"))
    product_tasks = json.loads((output / "product-tasks.json").read_text(encoding="utf-8"))

    assert exit_code == 0
    assert len(items) == 3
    assert all(item["status"] == "ready_for_review" for item in items)
    assert product_tasks[0]["status"] == "ready_for_review"


def test_run_builds_reviewable_items_from_per_file_asset_manifest(tmp_path):
    products, rules, basic, search = write_sources(tmp_path)
    backend = tmp_path / "backend.csv"
    with backend.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=["商品ID", "目标坑位", "空坑位", "审核状态完整", "采集时间", "证据"],
        )
        writer.writeheader()
        writer.writerow(
            {
                "商品ID": "123",
                "目标坑位": "3",
                "空坑位": "1;2;3",
                "审核状态完整": "true",
                "采集时间": "2026-07-17T09:00:00+08:00",
                "证据": "fixture",
            }
        )

    asset_manifest = tmp_path / "assets.csv"
    with asset_manifest.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=["product_id", "sku", "source_path", "license_status", "asset_type"],
        )
        writer.writeheader()
        for index in range(9):
            image = tmp_path / f"manifest-{index}.png"
            Image.new("RGB", (400, 400), color=(index, index, index)).save(image)
            writer.writerow(
                {
                    "product_id": "123",
                    "sku": "SKU-1",
                    "source_path": str(image),
                    "license_status": "confirmed",
                    "asset_type": "image",
                }
            )

    copy_responses = tmp_path / "copy-responses.json"
    copy_responses.write_text(
        json.dumps(
            {
                f"123:{slot}": {
                    "title": f"KK树便携水杯{slot}",
                    "description": "便携水杯设计，满足日常携带和饮水使用需求。",
                }
                for slot in (1, 2, 3)
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    output = tmp_path / "manifest-run"

    exit_code = main(
        [
            "run",
            "--month", "7",
            "--store", "KK Tree",
            "--products", str(products),
            "--rules", str(rules),
            "--basic", str(basic),
            "--search", str(search),
            "--backend-status", str(backend),
            "--asset-manifest", str(asset_manifest),
            "--copy-responses", str(copy_responses),
            "--output", str(output),
            "--started-at", "2026-07-17T10:00:00+08:00",
        ]
    )
    items = json.loads((output / "material-items.json").read_text(encoding="utf-8"))

    assert exit_code == 0
    assert len(items) == 3
    assert all(item["status"] == "ready_for_review" for item in items)
    assert all(
        asset["license_status"] == "confirmed"
        for item in items
        for asset in item["assets"]
    )
    assert all(
        asset["source_system"] == "manifest"
        for item in items
        for asset in item["assets"]
    )


def test_run_rejects_directory_and_manifest_asset_sources_together():
    with pytest.raises(SystemExit):
        build_parser().parse_args(
            [
                "run",
                "--month", "7",
                "--store", "KK Tree",
                "--products", "products.csv",
                "--rules", "rules.csv",
                "--basic", "basic.xlsx",
                "--search", "search.xlsx",
                "--asset-root", "assets",
                "--asset-manifest", "assets.csv",
                "--output", "output",
            ]
        )


def test_page_context_connects_to_runtime_cdp_without_storing_session():
    sentinel = object()
    calls = []

    @contextmanager
    def factory(url):
        calls.append(url)
        yield sentinel

    with page_context(None, "http://127.0.0.1:9222", factory) as resolved:
        assert resolved is sentinel

    assert calls == ["http://127.0.0.1:9222"]


def test_page_context_requires_injected_page_or_cdp_url():
    with pytest.raises(BrowserSessionRequired):
        with page_context(None, None, None):
            pass


def test_export_command_writes_isolated_manifest_for_basic_and_promotion(tmp_path):
    basic = tmp_path / "basic-source.xlsx"
    search = tmp_path / "search-source.xlsx"
    basic_workbook = Workbook()
    basic_sheet = basic_workbook.active
    basic_sheet.append(["商品ID", "商品标题", "商品白底图", "短标题"])
    basic_sheet.append(["123", "儿童帽", "", ""])
    basic_workbook.save(basic)
    search_workbook = Workbook()
    search_sheet = search_workbook.active
    search_sheet.append(["商品ID", "素材类型", "素材ID", "审核状态"])
    search_sheet.append(["123", "图文", "M-1", "审核通过"])
    search_workbook.save(search)
    page = CliFakePage(
        [CliFakeDownload(basic, "基础素材.xlsx"), CliFakeDownload(search, "推广素材.xlsx")]
    )
    selectors = Path(__file__).parents[1] / "config" / "selectors.example.yaml"
    output = tmp_path / "exports"

    exit_code = main(
        [
            "export",
            "--store", "KK Tree",
            "--selectors", str(selectors),
            "--output", str(output),
            "--run-id", "RUN-1",
            "--downloaded-at", "2026-07-17T10:00:00+08:00",
            "--product-status", "售卖中",
            "--report", "both",
        ],
        page=page,
    )
    records = json.loads((output / "source-files.json").read_text(encoding="utf-8"))
    manifest = json.loads((output / "export-manifest.json").read_text(encoding="utf-8"))

    assert exit_code == 0
    assert len(records) == 2
    assert all(len(record["sha256"]) == 64 for record in records)
    assert [record["report_type"] for record in records] == ["basic", "promotion"]
    assert manifest["schema_version"] == 1
    assert manifest["run_id"] == "RUN-1"
    assert manifest["store"] == "KK Tree"
    assert manifest["filters"] == {"product_status": "售卖中"}
    assert manifest["status"] == "complete"
    assert manifest["promotion_contract_status"] == "draft"
    assert manifest["reports"] == records


def test_export_command_rejects_nonempty_output_before_browser_use(tmp_path):
    selectors = Path(__file__).parents[1] / "config" / "selectors.example.yaml"
    output = tmp_path / "exports"
    output.mkdir()
    sentinel = output / "keep.txt"
    sentinel.write_text("keep", encoding="utf-8")
    page = CliFakePage([])

    exit_code = main(
        [
            "export",
            "--store", "KK Tree",
            "--selectors", str(selectors),
            "--output", str(output),
            "--run-id", "RUN-1",
            "--downloaded-at", "2026-07-17T10:00:00+08:00",
            "--product-status", "售卖中",
            "--report", "basic",
        ],
        page=page,
    )

    assert exit_code == 2
    assert sentinel.read_text(encoding="utf-8") == "keep"
    assert page.clicked == []


def test_supplement_command_writes_backend_status_csv(tmp_path):
    candidates = tmp_path / "candidates.csv"
    with candidates.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["商品ID", "reason"])
        writer.writeheader()
        writer.writerow({"商品ID": "123", "reason": "BROWSER_SUPPLEMENT_REQUIRED"})
    page = CliFakePage()
    page.texts.update({"#product-id": "123", "#desired-slots": "3坑", "#material-table": "table"})
    page.visible.add("#material-table")
    page.counts[".material-row"] = 2
    page.lists[".empty-slot"] = ["坑位3"]
    page.lists[".review-status"] = ["审核通过", "审核中"]
    selectors = Path(__file__).parents[1] / "config" / "selectors.example.yaml"
    output = tmp_path / "backend.csv"

    exit_code = main(
        [
            "supplement",
            "--store", "KK Tree",
            "--selectors", str(selectors),
            "--candidates", str(candidates),
            "--output", str(output),
            "--collected-at", "2026-07-17T10:00:00+08:00",
        ],
        page=page,
    )
    with output.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))

    assert exit_code == 0
    assert rows[0]["商品ID"] == "123"
    assert rows[0]["空坑位"] == "3"


def test_supplement_defaults_to_high_value_batch_scan_and_writes_checkpoint(
    tmp_path,
    monkeypatch,
):
    import upload_search_materials.cli as cli_module

    first = {
        "商品ID": "565628742471",
        "目标容量": "9",
        "目标坑位": "9",
        "现有素材数": "9",
        "缺失数量": "0",
        "空坑位": "",
        "远端素材ID": "A;B",
        "素材状态": "",
        "审核状态": "",
        "状态完整": "false",
        "审核状态完整": "false",
        "状态": "ready_for_review",
        "原因码": "",
        "采集时间": "2026-07-24T01:57:30+08:00",
        "证据": "page=1",
    }
    second = dict(first, 商品ID="903588197784", 现有素材数="3", 缺失数量="6")

    def fake_scan(page, selectors, **kwargs):
        assert kwargs["filter_selector_key"] == "high_value_filter"
        kwargs["on_page"](1, [first])
        kwargs["on_page"](2, [first, second])
        return [first, second]

    monkeypatch.setattr(cli_module, "scan_recommended_material_status", fake_scan)
    page = CliFakePage()
    selectors = Path(__file__).parents[1] / "config" / "selectors.example.yaml"
    output = tmp_path / "promotion-material-status.csv"

    exit_code = main(
        [
            "supplement",
            "--store", "KK Tree",
            "--selectors", str(selectors),
            "--output", str(output),
            "--collected-at", "2026-07-24T01:57:30+08:00",
        ],
        page=page,
    )

    with output.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    checkpoint = json.loads(
        output.with_suffix(".checkpoint.json").read_text(encoding="utf-8")
    )

    assert exit_code == 0
    assert [row["商品ID"] for row in rows] == ["565628742471", "903588197784"]
    assert rows[1]["缺失数量"] == "6"
    assert checkpoint["status"] == "complete"
    assert checkpoint["last_completed_page"] == 2
    assert checkpoint["row_count"] == 2
    assert checkpoint["scan_mode"] == "high-value"


def test_resume_partitions_uncertain_item_into_verification_only(tmp_path):
    item = material_item_from_dict(
        {
            "task_id": "MAT-1",
            "product_id": "123",
            "material_type": "image_text",
            "slot_index": 1,
            "status": "approved",
            "assets": [],
            "title": "标题",
            "description": "描述内容满足长度要求。",
        }
    )
    state = StateStore(tmp_path / "run.sqlite3")
    state.save_item("MAT-1", "publish_uncertain", evidence="timeout", attempt_count=1)

    upload_items, verification_items = partition_persisted_items(state, [item], resume=True)

    assert upload_items == []
    assert verification_items == [item]
    state.close()


def test_repeated_publish_does_not_requeue_submitted_item(tmp_path):
    item = material_item_from_dict(
        {
            "task_id": "MAT-1",
            "product_id": "123",
            "material_type": "image_text",
            "slot_index": 1,
            "status": "approved",
            "assets": [],
            "title": "标题",
            "description": "描述内容满足长度要求。",
        }
    )
    state = StateStore(tmp_path / "run.sqlite3")
    state.save_item("MAT-1", "submitted", remote_material_id="RM-1", evidence="remote row")

    upload_items, verification_items = partition_persisted_items(state, [item], resume=False)

    assert upload_items == []
    assert verification_items == []
    state.close()


def test_publish_uncertain_pauses_remaining_batch_items(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    items = [
        {
            "task_id": task_id,
            "product_id": product_id,
            "material_type": "image_text",
            "slot_index": 1,
            "status": "approved",
            "assets": [],
            "title": "标题",
            "description": "描述内容满足长度要求。",
        }
        for task_id, product_id in (("MAT-1", "123"), ("MAT-2", "456"))
    ]
    (run_dir / "material-items.json").write_text(json.dumps(items), encoding="utf-8")
    (run_dir / "run.json").write_text(
        json.dumps(
            {
                "run_id": "RUN-1",
                "store": "KK Tree",
                "source_files": {},
                "source_sha256": {},
            }
        ),
        encoding="utf-8",
    )
    manifest_items = [material_item_from_dict(value) for value in items]
    (run_dir / "approval-manifest.json").write_text(
        json.dumps(
            create_manifest(
                "KK Tree",
                manifest_items,
                "operator",
                "2026-07-17T10:00:00+08:00",
                valid_until="2099-07-18T10:00:00+08:00",
                run_id="RUN-1",
                source_sha256={},
            )
        ),
        encoding="utf-8",
    )
    calls = []

    def fake_upload(page, item, manifest, selectors, **kwargs):
        calls.append(item.task_id)
        kwargs["before_publish"]()
        return UploadOutcome("publish_uncertain", "PUBLISH_UNCERTAIN", evidence="timeout")

    monkeypatch.setattr(cli_module, "upload_approved_item", fake_upload)
    selectors = Path(__file__).parents[1] / "config" / "selectors.example.yaml"

    missing_state_exit = main(
        [
            "publish",
            "--run-dir", str(run_dir),
            "--store", "KK Tree",
            "--selectors", str(selectors),
        ],
        page=object(),
    )
    assert missing_state_exit == 2
    assert calls == []
    assert not (run_dir / "run.sqlite3").exists()

    state = StateStore(run_dir / "run.sqlite3")
    state.save_item("MAT-1", "approved")
    state.save_item("MAT-2", "approved")
    state.close()

    exit_code = main(
        [
            "publish",
            "--run-dir", str(run_dir),
            "--store", "KK Tree",
            "--selectors", str(selectors),
        ],
        page=object(),
    )

    assert exit_code == 1
    assert calls == ["MAT-1"]
    first_results = (run_dir / "upload-results.json").read_text(encoding="utf-8")

    repeated_exit = main(
        [
            "publish",
            "--run-dir", str(run_dir),
            "--store", "KK Tree",
            "--selectors", str(selectors),
        ],
        page=object(),
    )

    assert repeated_exit == 2
    assert calls == ["MAT-1"]
    assert (run_dir / "upload-results.json").read_text(encoding="utf-8") == first_results


def _write_index_products(path: Path, rows=(("123", "SKU-1", "Hat"),)):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=PRODUCT_HEADERS)
        writer.writeheader()
        for product_id, sku, title in rows:
            writer.writerow(
                {
                    "商品ID": product_id,
                    "商品名称（查找引用）": title,
                    "货号（查找引用）": sku,
                    "产品等级": "A",
                    "链接": "https://example.invalid/product",
                    "运营": "owner",
                    "组别": "team",
                    "品类-公司维度划分": "category",
                }
            )


def _index_assets_args(products: Path, root: Path, output: Path, *extra: str):
    return [
        "index-assets",
        "--products", str(products),
        "--root", f"model={root}",
        "--output", str(output),
        *extra,
    ]


def test_index_assets_cli_writes_sqlite_csv_json_and_never_uses_browser(
    tmp_path, monkeypatch
):
    products = tmp_path / "products.csv"
    _write_index_products(products)
    root = tmp_path / "assets"
    image_dir = root / "catalog" / "123"
    image_dir.mkdir(parents=True)
    Image.new("RGB", (400, 400), color="white").save(image_dir / "a.png")
    output = tmp_path / "index-output"
    boundary_calls = []

    def forbidden_boundary(*args, **kwargs):
        boundary_calls.append((args, kwargs))
        pytest.fail("index-assets must not enter publish or browser boundaries")

    monkeypatch.setattr(cli_module, "_publish", forbidden_boundary)
    monkeypatch.setattr(cli_module, "open_cdp_page", forbidden_boundary)

    exit_code = main(_index_assets_args(products, root, output), page=object())

    assert exit_code == 0
    assert sorted(path.name for path in output.iterdir()) == [
        "asset-index.sqlite3",
        "match-candidates.csv",
        "scan-summary.json",
    ]
    summary = json.loads((output / "scan-summary.json").read_text(encoding="utf-8"))
    assert summary["complete"] is True
    assert summary["matched"] == 1
    assert boundary_calls == []


def test_index_assets_cli_binds_confirmed_root_from_interaction_input(tmp_path):
    products = tmp_path / "products.csv"
    _write_index_products(products)
    root = tmp_path / "confirmed-product-folder"
    root.mkdir()
    Image.new("RGB", (400, 400), color="white").save(root / "a.png")
    decisions = tmp_path / "input.json"
    decisions.write_text(
        json.dumps(
            {
                "values": {
                    "folder_decisions": [
                        {
                            "decision": "confirmed",
                            "folder_path": str(root.resolve()),
                            "product_id": "123",
                        }
                    ]
                }
            }
        ),
        encoding="utf-8",
    )
    output = tmp_path / "index-output"

    exit_code = main(
        _index_assets_args(
            products,
            root,
            output,
            "--folder-decisions",
            str(decisions),
        )
    )

    assert exit_code == 0
    summary = json.loads(
        (output / "scan-summary.json").read_text(encoding="utf-8")
    )
    assert summary["matched"] == 1
    assert summary["confirmed_folder_binding"] == {
        "bound_roots": 1,
        "bound_files": 1,
        "inspection_failures": 0,
    }
    with (output / "match-candidates.csv").open(
        "r", encoding="utf-8-sig", newline=""
    ) as stream:
        rows = list(csv.DictReader(stream))
    assert rows[0]["product_id"] == "123"
    assert rows[0]["hashes_complete"] == "true"


def test_prepare_gallery_cli_reads_index_and_material_gap_without_browser(
    tmp_path,
):
    products = tmp_path / "products.csv"
    _write_index_products(products)
    root = tmp_path / "assets"
    image_dir = root / "catalog" / "123"
    image_dir.mkdir(parents=True)
    Image.new("RGB", (300, 400), color="white").save(image_dir / "a.png")
    index_output = tmp_path / "index-output"
    assert main(_index_assets_args(products, root, index_output)) == 0
    status_path = tmp_path / "promotion-status.csv"
    with status_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=["商品ID", "缺失数量"],
        )
        writer.writeheader()
        writer.writerow({"商品ID": "123", "缺失数量": "2"})
    gallery_path = tmp_path / "asset-gallery.json"

    exit_code = main(
        [
            "prepare-gallery",
            "--index",
            str(index_output / "asset-index.sqlite3"),
            "--status",
            str(status_path),
            "--output",
            str(gallery_path),
        ],
        page=object(),
    )

    gallery = json.loads(gallery_path.read_text(encoding="utf-8"))
    assert exit_code == 0
    assert gallery["selection_mode"] == "deterministic"
    assert gallery["remote_dedupe_status"] == "not_available"
    assert gallery["requirements"] == [
        {
            "product_id": "123",
            "product_title": "Hat",
            "missing_materials": 2,
            "slot_image_min": 3,
            "slot_image_max": 9,
            "slot_planning_stage": "slots_copy",
        }
    ]
    assert gallery["asset_candidates"][0]["source_path"] == str(
        (image_dir / "a.png").resolve()
    )
    assert gallery["asset_candidates"][0]["license_status"] == "unknown"


def test_index_assets_cli_refresh_interrupt_executes_resume_argv_and_keeps_checkpoint(
    tmp_path, monkeypatch, capsys
):
    products = tmp_path / "products.csv"
    _write_index_products(products)
    root = tmp_path / "assets"
    for directory in ("a", "b"):
        image_dir = root / directory / "123"
        image_dir.mkdir(parents=True)
        Image.new("RGB", (400, 400), color="white").save(image_dir / f"{directory}.png")
    output = tmp_path / "index-output"
    args = _index_assets_args(
        products, root, output, "--partition-depth", "1", "--checkpoint-size", "1"
    )
    assert main(args) == 0
    capsys.readouterr()
    real_restart = cli_module.AssetIndexStore.restart_partition
    interrupted = False

    def interrupt_before_second_partition(self, partition_id, scan_id):
        nonlocal interrupted
        partition = self._partition(partition_id)
        if partition["relative_path"] == "b" and not interrupted:
            interrupted = True
            raise KeyboardInterrupt
        return real_restart(self, partition_id, scan_id)

    monkeypatch.setattr(
        cli_module.AssetIndexStore,
        "restart_partition",
        interrupt_before_second_partition,
    )
    interrupted_exit = main([*args, "--refresh"])
    interrupted_stderr = capsys.readouterr().err
    with sqlite3.connect(output / "asset-index.sqlite3") as connection:
        mixed_scan_ids = {
            row[0] for row in connection.execute("SELECT current_scan_id FROM partitions")
        }
        active_before_resume = connection.execute(
            "SELECT value FROM scan_meta WHERE key='active_scan_id'"
        ).fetchone()
    summary = json.loads((output / "scan-summary.json").read_text(encoding="utf-8"))

    resumed_exit = main([*args, "--resume"])

    assert interrupted_exit == 1
    assert "checkpoint" in interrupted_stderr and "--resume" in interrupted_stderr
    assert summary["resume_command"].count("--root") == 1
    assert len(mixed_scan_ids) == 2
    assert active_before_resume is not None
    assert resumed_exit == 0
    assert sorted(path.name for path in output.iterdir()) == [
        "asset-index.sqlite3", "match-candidates.csv", "scan-summary.json",
    ]
    with sqlite3.connect(output / "asset-index.sqlite3") as connection:
        assert connection.execute(
            "SELECT value FROM scan_meta WHERE key='active_scan_id'"
        ).fetchone() is None


def test_index_assets_cli_successful_resume_and_refresh_complete_chains(tmp_path):
    products = tmp_path / "products.csv"
    _write_index_products(products)
    root = tmp_path / "assets"
    image_dir = root / "catalog" / "123"
    image_dir.mkdir(parents=True)
    Image.new("RGB", (400, 400), color="white").save(image_dir / "a.png")
    output = tmp_path / "index-output"
    args = _index_assets_args(products, root, output)

    assert main(args) == 0
    assert main([*args, "--resume"]) == 0
    Image.new("RGB", (400, 400), color="black").save(image_dir / "b.png")
    assert main([*args, "--refresh"]) == 0

    summary = json.loads((output / "scan-summary.json").read_text(encoding="utf-8"))
    assert summary["mode"] == "refresh"
    assert summary["matched"] == 2


def test_index_assets_cli_keeps_row_validation_local_and_persists_evidence(tmp_path):
    products = tmp_path / "products.csv"
    rows = [("123", "VALID-SKU", "Valid Hat")]
    rows.extend(("", f"MISSING-{index}", f"Missing {index}") for index in range(33))
    rows.extend(("999", f"DUP-{index}", f"Duplicate {index}") for index in range(11))
    _write_index_products(products, rows)
    root = tmp_path / "assets"
    for directory, name in (
        ("123", "valid.png"),
        ("999", "duplicate.png"),
        ("MISSING-0", "missing.png"),
    ):
        target = root / directory
        target.mkdir(parents=True)
        Image.new("RGB", (40, 40), color="white").save(target / name)
    output = tmp_path / "index-output"

    exit_code = main(_index_assets_args(products, root, output))

    assert exit_code == 0
    with (output / "match-candidates.csv").open(
        "r", encoding="utf-8-sig", newline=""
    ) as stream:
        candidates = list(csv.DictReader(stream))
    assert {(row["product_id"], row["image_count"]) for row in candidates} == {
        ("123", "1")
    }
    summary = json.loads((output / "scan-summary.json").read_text(encoding="utf-8"))
    validation = summary["product_validation"]
    assert validation == {
        "row_count": 45,
        "eligible_count": 1,
        "blocked_count": 44,
        "reason_code_counts": {
            "DUPLICATE_PRODUCT_ID": 11,
            "MISSING_PRODUCT_ID": 33,
        },
        "blocked_rows": [
            {
                "source_row": source_row,
                "reason_codes": [
                    "MISSING_PRODUCT_ID"
                    if source_row <= 35
                    else "DUPLICATE_PRODUCT_ID"
                ],
            }
            for source_row in range(3, 47)
        ],
    }
    assert summary["database"]["active_files"] == 3
    assert summary["matched"] == 1


def test_index_assets_cli_all_row_blocked_table_still_builds_metadata_index(tmp_path):
    products = tmp_path / "products.csv"
    _write_index_products(
        products,
        (("", "MISSING-SKU", "Missing ID"), ("bad-id", "BAD-SKU", "Invalid ID")),
    )
    root = tmp_path / "assets"
    root.mkdir()
    Image.new("RGB", (40, 40), color="white").save(root / "plain.png")
    output = tmp_path / "index-output"

    assert main(_index_assets_args(products, root, output)) == 0

    summary = json.loads((output / "scan-summary.json").read_text(encoding="utf-8"))
    assert summary["product_validation"]["eligible_count"] == 0
    assert summary["product_validation"]["blocked_count"] == 2
    assert summary["database"]["active_files"] == 1
    assert summary["matched"] == 0


def test_index_assets_cli_empty_product_table_is_batch_blocking(tmp_path):
    products = tmp_path / "products.csv"
    _write_index_products(products, ())
    root = tmp_path / "assets"
    root.mkdir()
    output = tmp_path / "index-output"

    assert main(_index_assets_args(products, root, output)) == 2
    assert not output.exists()


def test_index_assets_new_rejects_any_nonempty_output_without_touching_sentinels(
    tmp_path,
):
    products = tmp_path / "products.csv"
    _write_index_products(products)
    root = tmp_path / "assets"
    root.mkdir()
    output = tmp_path / "index-output"
    output.mkdir()
    sentinel = output / "keep.bin"
    candidates = output / "match-candidates.csv"
    summary = output / "scan-summary.json"
    sentinel.write_bytes(b"\x00keep-me\xff")
    candidates.write_bytes(b"existing,csv\r\n")
    summary.write_bytes(b'{"existing":true}\n')
    (output / "existing-directory").mkdir()
    before = {
        path.name: path.read_bytes()
        for path in (sentinel, candidates, summary)
    }

    assert main(_index_assets_args(products, root, output)) == 2

    assert {
        path.name: path.read_bytes()
        for path in (sentinel, candidates, summary)
    } == before
    assert (output / "existing-directory").is_dir()
    assert not (output / "asset-index.sqlite3").exists()


def test_index_assets_cli_rejects_declared_symlink_root_before_output(tmp_path):
    products = tmp_path / "products.csv"
    _write_index_products(products)
    target = tmp_path / "assets"
    target.mkdir()
    root = tmp_path / "linked-assets"
    try:
        root.symlink_to(target, target_is_directory=True)
    except OSError as error:
        pytest.skip(f"symlink creation is not permitted: {error}")
    output = tmp_path / "index-output"

    assert main(_index_assets_args(products, root, output)) == 2
    assert not output.exists()


def test_index_assets_cli_rejects_declared_reparse_root_before_output(
    tmp_path, monkeypatch
):
    products = tmp_path / "products.csv"
    _write_index_products(products)
    root = tmp_path / "assets"
    root.mkdir()
    output = tmp_path / "index-output"
    real_lstat = Path.lstat

    def reparse_lstat(path):
        if Path(path) == root:
            return SimpleNamespace(
                st_mode=stat.S_IFDIR,
                st_file_attributes=getattr(
                    stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0x400
                ),
            )
        return real_lstat(path)

    monkeypatch.setattr(Path, "lstat", reparse_lstat)

    assert main(_index_assets_args(products, root, output)) == 2
    assert not output.exists()


@pytest.mark.parametrize(
    "option,value",
    [("--partition-depth", "0"), ("--checkpoint-size", "0")],
)
def test_index_assets_positive_options_are_validated_before_output_or_database(
    tmp_path, option, value
):
    products = tmp_path / "products.csv"
    _write_index_products(products)
    root = tmp_path / "assets"
    root.mkdir()
    output = tmp_path / "index-output"

    assert main(_index_assets_args(products, root, output, option, value)) == 2
    assert not output.exists()
    assert not (output / "asset-index.sqlite3").exists()


def test_index_assets_empty_root_path_returns_two_without_creating_output(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    products = tmp_path / "products.csv"
    _write_index_products(products)
    output = tmp_path / "index-output"

    exit_code = main(
        [
            "index-assets", "--products", str(products), "--root", "model=",
            "--output", str(output),
        ]
    )

    assert exit_code == 2
    assert not output.exists()


@pytest.mark.parametrize("case", ["duplicate-source", "missing-root", "existing-without-mode"])
def test_index_assets_cli_rejects_invalid_roots_and_existing_database(tmp_path, case):
    products = tmp_path / "products.csv"
    _write_index_products(products)
    root = tmp_path / "assets"
    root.mkdir()
    output = tmp_path / "index-output"
    args = _index_assets_args(products, root, output)
    if case == "duplicate-source":
        args[args.index("--output"):args.index("--output")] = ["--root", f"model={root}"]
    elif case == "missing-root":
        args[args.index(f"model={root}")] = f"model={tmp_path / 'missing'}"
    else:
        assert main(args) == 0

    assert main(args) == 2


def test_index_assets_cli_rejects_resume_refresh_together():
    with pytest.raises(SystemExit) as error:
        main(
            [
                "index-assets", "--products", "products.csv", "--root", "model=assets",
                "--output", "output", "--resume", "--refresh",
            ]
        )
    assert error.value.code == 2


@pytest.mark.parametrize("corruption", ["identity", "schema"])
def test_index_assets_cli_rejects_identity_or_schema_mismatch(tmp_path, corruption):
    products = tmp_path / "products.csv"
    _write_index_products(products)
    root = tmp_path / "assets"
    root.mkdir()
    output = tmp_path / "index-output"
    args = _index_assets_args(products, root, output)
    assert main(args) == 0
    database = output / "asset-index.sqlite3"
    if corruption == "identity":
        _write_index_products(products, (("123", "SKU-1", "Changed Hat"),))
    else:
        connection = sqlite3.connect(database)
        connection.execute("UPDATE scan_meta SET value='999' WHERE key='schema_version'")
        connection.commit()
        connection.close()

    assert main([*args, "--resume"]) == 2


def test_index_assets_cli_returns_one_for_partial_partition_failure(tmp_path, monkeypatch):
    products = tmp_path / "products.csv"
    _write_index_products(products)
    root = tmp_path / "assets"
    (root / "good" / "123").mkdir(parents=True)
    (root / "bad").mkdir()
    Image.new("RGB", (400, 400), color="white").save(root / "good" / "123" / "a.png")
    output = tmp_path / "index-output"
    real_scandir = cli_module.os.scandir

    def fail_bad(path):
        if Path(path) == (root / "bad").resolve():
            raise OSError("partition unavailable")
        return real_scandir(path)

    monkeypatch.setattr("upload_search_materials.asset_index.os.scandir", fail_bad)

    assert main(_index_assets_args(products, root, output)) == 1
    summary = json.loads((output / "scan-summary.json").read_text(encoding="utf-8"))
    assert summary["partial_failure"] is True
    assert summary["failed"] == 1
    assert summary["errors"][0]["code"] == "PARTITION_ENUMERATION_ERROR"
    assert summary["errors"][0]["detail"] == "partition unavailable"


def test_index_assets_help_lists_all_options(capsys):
    with pytest.raises(SystemExit) as error:
        build_parser().parse_args(["index-assets", "--help"])
    assert error.value.code == 0
    help_text = capsys.readouterr().out
    for option in (
        "--products", "--root", "--output", "--partition-depth", "--checkpoint-size",
        "--resume", "--refresh",
    ):
        assert option in help_text
