from contextlib import contextmanager
import hashlib
import json
from pathlib import Path

import pytest

from upload_search_materials.browser.config import SelectorConfigError, load_selectors
from upload_search_materials.browser.export_page import export_reports
from upload_search_materials.browser.material_page import (
    PaginationStateError,
    SelectorInvalidError,
    _click_with_popup_retries,
    _settle_safe_popups,
    scan_recommended_material_status,
    supplement_material_status,
)
from upload_search_materials.browser.session import (
    HumanCheckRequired,
    StoreIdentityError,
    assert_store_identity,
    detect_human_check,
)
from upload_search_materials.browser.upload_page import upload_approved_item
from upload_search_materials.browser.verifier import verify_remote_item
from upload_search_materials.approval import create_manifest
from upload_search_materials.models import AssetRecord, MaterialItem, MaterialStatus


def test_false_success_fixture_matches_previous_final_page():
    fixture = json.loads(
        (
            Path(__file__).parent
            / "fixtures"
            / "session_20260730_032916_pagination_origin_sanitized.json"
        ).read_text(encoding="utf-8")
    )

    assert fixture["preceding_collection"]["row_count"] == 255
    assert fixture["preceding_collection"]["page_count"] == 26
    assert (
        fixture["preceding_collection"]["final_page_product_ids"]
        == fixture["false_success_collection"]["all_product_ids"]
    )
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError


REQUIRED_SELECTOR_VALUES = {
    "store_name": "#store",
    "human_check": "#human-check",
    "export_basic": "#export-basic",
    "export_search": "#export-search",
    "product_search": "#product-search",
    "product_id": "#product-id",
    "desired_slots": "#desired-slots",
    "material_table": "#material-table",
    "material_rows": ".material-row",
    "empty_slots": ".empty-slot",
    "review_status": ".review-status",
    "image_text_action": "#image-text-action",
    "video_action": "#video-action",
    "file_input": "#file-input",
    "title_input": "#title-input",
    "description_input": "#description-input",
    "publish_button": "#publish",
    "success_signal": "#success",
    "remote_material_id": "#remote-id",
    "remote_material_table": "#remote-table",
    "remote_material_fingerprint": "#remote-fingerprint",
    "remote_material_status": "#remote-status",
    "remote_material_slot": "#remote-slot",
    "remote_material_time": "#remote-time",
}


class FakeLocator:
    def __init__(self, page, selector):
        self.page = page
        self.selector = selector

    def inner_text(self):
        return self.page.texts.get(self.selector, "")

    def all_inner_texts(self):
        return list(self.page.lists.get(self.selector, []))

    def count(self):
        if self.selector in self.page.counts:
            return self.page.counts[self.selector]
        return 1 if self.selector in self.page.texts else 0

    def is_visible(self):
        return self.selector in self.page.visible

    def fill(self, value):
        self.page.filled.append((self.selector, value))

    def set_input_files(self, values):
        self.page.input_files.append((self.selector, list(values)))

    def press(self, key):
        self.page.pressed.append((self.selector, key))

    def click(self):
        self.page.clicked.append(self.selector)
        error = self.page.click_errors.get(self.selector)
        if error:
            raise error


class FakePage:
    def __init__(self):
        self.texts = {}
        self.lists = {}
        self.counts = {}
        self.visible = set()
        self.filled = []
        self.pressed = []
        self.clicked = []
        self.input_files = []
        self.click_errors = {}

    def locator(self, selector):
        return FakeLocator(self, selector)


class FakeDownload:
    def __init__(self, source, suggested_filename):
        self.source = Path(source)
        self.suggested_filename = suggested_filename

    def save_as(self, destination):
        Path(destination).write_bytes(self.source.read_bytes())


class FakeDownloadInfo:
    def __init__(self, download):
        self.value = download


class FakeExportPage(FakePage):
    def __init__(self, downloads):
        super().__init__()
        self.downloads = iter(downloads)

    @contextmanager
    def expect_download(self):
        yield FakeDownloadInfo(next(self.downloads))


def write_selectors(path, values):
    import yaml

    path.write_text(yaml.safe_dump(values, allow_unicode=True), encoding="utf-8")


def test_missing_required_selector_fails_configuration(tmp_path):
    path = tmp_path / "selectors.yaml"
    values = dict(REQUIRED_SELECTOR_VALUES)
    del values["publish_button"]
    write_selectors(path, values)

    with pytest.raises(SelectorConfigError, match="publish_button"):
        load_selectors(path)


def test_promotion_export_selector_accepts_new_name_and_normalizes_legacy_alias(
    tmp_path,
):
    path = tmp_path / "selectors.yaml"
    values = dict(REQUIRED_SELECTOR_VALUES)
    values["export_promotion"] = values.pop("export_search")
    write_selectors(path, values)

    selectors = load_selectors(path)

    assert selectors["export_promotion"] == "#export-search"
    assert selectors["export_search"] == "#export-search"


def test_missing_both_promotion_export_selector_names_fails_configuration(tmp_path):
    path = tmp_path / "selectors.yaml"
    values = dict(REQUIRED_SELECTOR_VALUES)
    values.pop("export_search")
    write_selectors(path, values)

    with pytest.raises(SelectorConfigError, match="export_promotion"):
        load_selectors(path)


def test_wrong_store_stops_batch():
    page = FakePage()
    page.texts["#store"] = "Other Store"

    with pytest.raises(StoreIdentityError):
        assert_store_identity(page, "#store", "KK Tree")


def test_human_check_stops_batch():
    page = FakePage()
    page.visible.add("#human-check")

    with pytest.raises(HumanCheckRequired):
        detect_human_check(page, "#human-check")


def test_missing_material_table_is_not_interpreted_as_zero():
    page = FakePage()
    page.texts["#product-id"] = "123"
    page.texts["#desired-slots"] = "3坑"
    selectors = dict(REQUIRED_SELECTOR_VALUES)

    rows = supplement_material_status(
        page,
        selectors,
        ["123"],
        collected_at="2026-07-17T10:00:00+08:00",
    )

    assert rows[0]["状态"] == "needs_manual_review"
    assert rows[0]["原因码"] == "SELECTOR_INVALID"
    assert rows[0]["现有素材数"] == ""
    assert rows[0]["空坑位"] == ""
    assert "failed_field=material_table" in rows[0]["证据"]


def test_material_status_supplement_reads_exact_product():
    page = FakePage()
    page.texts.update({"#product-id": "123", "#desired-slots": "3坑", "#material-table": "table"})
    page.visible.add("#material-table")
    page.counts[".material-row"] = 2
    page.lists[".empty-slot"] = ["坑位3"]
    page.lists[".review-status"] = ["审核通过", "审核中"]

    rows = supplement_material_status(
        page,
        REQUIRED_SELECTOR_VALUES,
        ["123"],
        collected_at="2026-07-17T10:00:00+08:00",
    )

    assert rows == [
        {
            "商品ID": "123",
            "目标坑位": "3",
                "现有素材数": "2",
                "空坑位": "3",
                "精确坑位状态": "collected",
                "审核状态": "审核通过;审核中",
            "审核状态完整": "true",
            "状态": "ready_for_review",
            "原因码": "",
            "采集时间": "2026-07-17T10:00:00+08:00",
            "证据": "product=123;rows=2;selector_version=runtime",
        }
    ]


class FakePromotionLocator:
    def __init__(self, page, selector):
        self.page = page
        self.selector = selector

    def click(self, **_kwargs):
        self.page.clicked.append(self.selector)
        if self.selector == "#next":
            self.page.page_index += 1
        elif self.selector == "#first":
            self.page.page_index = 0
        elif self.selector == "#recommended":
            self.page.recommended_checked = True

    def count(self):
        if self.selector == ".promotion-row":
            return len(self.page.pages[self.page.page_index])
        return 1

    def all_inner_texts(self):
        if self.selector == ".promotion-row":
            return self.page.pages[self.page.page_index]
        return []

    def get_attribute(self, name):
        if self.selector == "#recommended" and name == "aria-checked":
            return "true" if self.page.recommended_checked else "false"
        return None

    def inner_text(self):
        if self.selector == "#current":
            return str(self.page.page_index + 1)
        if self.selector == "#first":
            return "1"
        if self.selector == "#terminal":
            return str(len(self.page.pages))
        return ""

    def is_enabled(self):
        return self.page.page_index < len(self.page.pages) - 1


class FakePromotionPage:
    def __init__(self, pages, *, page_index=0):
        self.pages = pages
        self.page_index = page_index
        self.recommended_checked = False
        self.clicked = []
        self.waited = []

    def locator(self, selector):
        return FakePromotionLocator(self, selector)

    def wait_for_timeout(self, milliseconds):
        self.waited.append(milliseconds)


class DelayedPromotionLocator(FakePromotionLocator):
    def click(self, **_kwargs):
        if self.selector == "#next":
            self.page.clicked.append(self.selector)
            self.page.pending_next = True
            return
        super().click(**_kwargs)


class DelayedPromotionPage(FakePromotionPage):
    def __init__(self, pages):
        super().__init__(pages)
        self.pending_next = False

    def locator(self, selector):
        return DelayedPromotionLocator(self, selector)

    def wait_for_timeout(self, milliseconds):
        super().wait_for_timeout(milliseconds)
        if self.pending_next:
            self.page_index += 1
            self.pending_next = False


class FakePopupLocator:
    def __init__(self, page, selector):
        self.page = page
        self.selector = selector

    def count(self):
        return 1 if self.page.overlay_open else 0

    def nth(self, _index):
        return self

    def is_visible(self):
        return self.page.overlay_open

    def inner_text(self):
        return (
            f"guide-step-{self.page.guide_steps}"
            if self.selector == "#guide-next"
            else "popup-close"
        )

    def click(self, **_kwargs):
        self.page.clicked.append(self.selector)
        if self.selector == "#guide-next":
            self.page.guide_steps -= 1
            if self.page.guide_steps == 0:
                self.page.overlay_open = False
            return
        if self.page.click_races and self.selector == "#opened-overlay-close":
            self.page.click_races -= 1
            raise PlaywrightTimeoutError("popup rerendered")
        if self.selector == "#opened-overlay-close":
            self.page.overlay_open = False


class FakePopupPage:
    def __init__(self, *, click_races=0, guide_steps=0):
        self.overlay_open = True
        self.click_races = click_races
        self.guide_steps = guide_steps
        self.clicked = []
        self.waited = []

    def locator(self, selector):
        return FakePopupLocator(self, selector)

    def wait_for_timeout(self, milliseconds):
        self.waited.append(milliseconds)


def test_popup_settle_prioritizes_open_overlay_close_control():
    page = FakePopupPage()

    closed = _settle_safe_popups(
        page,
        {
            "safe_popup_close_priority": "#opened-overlay-close",
            "safe_popup_close": "#generic-close",
        },
        delay_ms=0,
    )

    assert closed == 1
    assert page.clicked == ["#opened-overlay-close"]
    assert page._tmall_collection_events[-1]["changed"] is True


def test_popup_settle_retries_when_overlay_rerenders_during_click():
    page = FakePopupPage(click_races=1)

    closed = _settle_safe_popups(
        page,
        {
            "safe_popup_close_priority": "#opened-overlay-close",
        },
        delay_ms=0,
    )

    assert closed == 1
    assert page.clicked == [
        "#opened-overlay-close",
        "#opened-overlay-close",
    ]


def test_popup_settle_advances_all_seven_scoped_guide_steps():
    page = FakePopupPage(guide_steps=7)

    closed = _settle_safe_popups(
        page,
        {
            "safe_popup_progress": "#guide-next",
            "safe_popup_close_priority": "#opened-overlay-close",
        },
        delay_ms=0,
    )

    assert closed == 7
    assert page.clicked == ["#guide-next"] * 7


def test_popup_settle_stops_one_noop_control_after_two_unchanged_actions():
    page = FakePopupPage()

    class NoOp(FakePopupLocator):
        def click(self, **_kwargs):
            self.page.clicked.append(self.selector)

    page.locator = lambda selector: NoOp(page, selector)

    closed = _settle_safe_popups(
        page,
        {"safe_popup_close_priority": "#noop"},
        delay_ms=0,
    )

    assert closed == 0
    assert page.clicked == ["#noop", "#noop"]


def test_popup_settle_catches_guide_that_appears_after_initial_quiet_check():
    page = FakePopupPage(guide_steps=1)
    page.overlay_open = False
    original_wait = page.wait_for_timeout

    def delayed_open(milliseconds):
        original_wait(milliseconds)
        if len(page.waited) == 1:
            page.overlay_open = True

    page.wait_for_timeout = delayed_open

    closed = _settle_safe_popups(
        page,
        {"safe_popup_progress": "#guide-next"},
        delay_ms=10,
    )

    assert closed == 1
    assert page.clicked == ["#guide-next"]


class OverlayBlockedTarget:
    def __init__(self):
        self.attempts = []
        self.evaluations = []

    def click(self, **kwargs):
        self.attempts.append(kwargs)
        if not kwargs.get("force"):
            raise PlaywrightTimeoutError("guide overlay intercepts pointer events")

    def evaluate(self, expression):
        self.evaluations.append(expression)


def test_successful_collection_click_does_not_scan_popup_controls():
    class DirectTarget:
        def __init__(self):
            self.attempts = []

        def click(self, **kwargs):
            self.attempts.append(kwargs)

    page = FakePopupPage()
    target = DirectTarget()

    _click_with_popup_retries(
        page,
        target,
        {
            "safe_popup_progress": "#guide-next",
            "safe_popup_close_priority": "#opened-overlay-close",
        },
        field_name="promotion_next_page",
        delay_ms=1000,
    )

    assert target.attempts == [{"timeout": 3000}]
    assert page.clicked == []
    assert page.waited == []


def test_read_only_collection_click_uses_dom_click_after_popup_retries():
    page = FakePopupPage()
    page.overlay_open = False
    target = OverlayBlockedTarget()

    _click_with_popup_retries(
        page,
        target,
        {"safe_popup_progress": "#guide-next"},
        field_name="high_value_filter",
        delay_ms=0,
    )

    assert len(target.attempts) == 2
    assert target.evaluations == ["element => element.click()"]
    assert page._tmall_collection_events == [
        {
            "action": "dom_click",
            "target_field": "high_value_filter",
            "reason": "recognized_guide_interceptor",
            "read_only": True,
        }
    ]


def test_unknown_overlay_never_uses_force_click():
    page = FakePopupPage()
    page.overlay_open = False
    target = OverlayBlockedTarget()

    with pytest.raises(SelectorInvalidError, match="popup_blocked"):
        _click_with_popup_retries(
            page,
            target,
            {},
            field_name="high_value_filter",
            delay_ms=0,
        )

    assert len(target.attempts) == 2


def test_write_target_never_uses_force_click_even_for_recognized_guide():
    page = FakePopupPage()
    page.overlay_open = False
    target = OverlayBlockedTarget()

    with pytest.raises(SelectorInvalidError, match="popup_blocked"):
        _click_with_popup_retries(
            page,
            target,
            {"safe_popup_progress": "#guide-next"},
            field_name="publish_button",
            delay_ms=0,
        )

    assert len(target.attempts) == 2


def test_recommended_promotion_scan_paginates_deduplicates_and_checkpoints():
    full_row = (
        "满坑商品 商品ID 565628742471 "
        "ID 3468934367960259 重复或图片有删除 "
        "ID 1539780903132488 ID 3295138149726183 ID 1787739833959372 "
        "ID 3212796032484488 ID 754653047484578 ID 506129500907 "
        "ID 461532917400 ID 461356707236 "
        "该商品为搜推高价值商品，已上调发布坑位到9篇，当前发布9篇。"
    )
    missing_row = (
        "缺口商品 商品ID 903588197784 "
        "ID 567934680336 ID 515805956420 ID 515848659407 "
        "该商品为搜推高价值商品，已上调发布坑位到9篇，当前发布3篇。"
    )
    unknown_target_row = "普通商品 商品ID 1037160921729 ID 557984331788"
    page = FakePromotionPage(
        [
            ["演示商品 商品ID xxxxxxxxxxxxxxx ID xxxxxxxxxxx1", full_row, missing_row],
            [missing_row, unknown_target_row],
        ]
    )
    checkpoints = []

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "recommended_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-07-24T01:57:30+08:00",
        on_page=lambda page_number, values: checkpoints.append(
            (page_number, [row["商品ID"] for row in values])
        ),
        settle_delay_ms=0,
    )

    assert page.clicked == ["#promotion", "#recommended", "#next"]
    assert [row["商品ID"] for row in rows] == [
        "565628742471",
        "903588197784",
        "1037160921729",
    ]
    assert rows[0]["目标容量"] == "9"
    assert rows[0]["现有素材数"] == "9"
    assert rows[0]["原因码"] == "REMOTE_MATERIAL_WARNING"
    assert rows[1]["缺失数量"] == "6"
    assert rows[1]["远端素材ID"] == "567934680336;515805956420;515848659407"
    assert rows[2]["目标容量"] == ""
    assert rows[2]["原因码"] == "TARGET_CAPACITY_NOT_EXPLICIT"
    assert checkpoints == [
        (1, ["565628742471", "903588197784"]),
        (2, ["565628742471", "903588197784", "1037160921729"]),
    ]


def test_promotion_scan_waits_until_next_page_product_ids_change():
    page = DelayedPromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇、当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇、当前发布 1 篇"],
        ]
    )

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-07-24T01:57:30+08:00",
        filter_selector_key="high_value_filter",
        settle_delay_ms=0,
        action_wait_ms=1000,
    )

    assert [row["商品ID"] for row in rows] == ["100", "200"]
    assert 250 in page.waited


@pytest.mark.parametrize("initial_page", [0, 1, 2])
def test_high_value_scan_normalizes_reused_tab_to_first_page(initial_page):
    page = FakePromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇，当前发布 1 篇"],
            ["商品三 商品ID 300 发布坑位到 9 篇，当前发布 2 篇"],
        ],
        page_index=initial_page,
    )
    events = []

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-07-30T10:00:00+08:00",
        filter_selector_key="high_value_filter",
        settle_delay_ms=0,
        action_wait_ms=1000,
        on_pagination_event=events.append,
    )

    assert [row["商品ID"] for row in rows] == ["100", "200", "300"]
    assert ("#first" in page.clicked) is (initial_page > 0)
    assert events[0]["event_type"] == "origin"
    assert events[0]["current_page"] == 1
    assert events[-1]["event_type"] == "terminal"
    assert events[-1]["current_page"] == 3


def test_high_value_scan_fails_when_first_page_reset_does_not_move():
    page = FakePromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇，当前发布 1 篇"],
        ],
        page_index=1,
    )
    original_locator = page.locator

    class NoOpFirst(FakePromotionLocator):
        def click(self, **_kwargs):
            if self.selector == "#first":
                self.page.clicked.append(self.selector)
                return
            super().click(**_kwargs)

    page.locator = lambda selector: NoOpFirst(page, selector)

    with pytest.raises(
        PaginationStateError, match="PAGINATION_ORIGIN_RESET_FAILED"
    ):
        scan_recommended_material_status(
            page,
            {
                "promotion_tab": "#promotion",
                "high_value_filter": "#recommended",
                "promotion_rows": ".promotion-row",
                "promotion_current_page": "#current",
                "promotion_first_page": "#first",
                "promotion_terminal_page": "#terminal",
                "promotion_next_page": "#next",
            },
            collected_at="2026-07-30T10:00:00+08:00",
            filter_selector_key="high_value_filter",
            settle_delay_ms=0,
            action_wait_ms=500,
        )

    page.locator = original_locator


def test_high_value_scan_waits_for_first_page_rows_to_replace_final_rows():
    class StaleRowsLocator(FakePromotionLocator):
        def click(self, **_kwargs):
            if self.selector == "#first":
                self.page.clicked.append(self.selector)
                self.page.stale_rows = list(
                    self.page.pages[self.page.page_index]
                )
                self.page.stale_reads = 2
                self.page.page_index = 0
                return
            super().click(**_kwargs)

        def all_inner_texts(self):
            if (
                self.selector == ".promotion-row"
                and self.page.stale_reads > 0
            ):
                return self.page.stale_rows
            return super().all_inner_texts()

    page = FakePromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇，当前发布 1 篇"],
            ["商品三 商品ID 300 发布坑位到 9 篇，当前发布 2 篇"],
        ],
        page_index=2,
    )
    page.stale_rows = []
    page.stale_reads = 0
    page.locator = lambda selector: StaleRowsLocator(page, selector)
    original_wait = page.wait_for_timeout

    def settle_rows(milliseconds):
        original_wait(milliseconds)
        if page.stale_reads > 0:
            page.stale_reads -= 1

    page.wait_for_timeout = settle_rows

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-07-30T10:00:00+08:00",
        filter_selector_key="high_value_filter",
        settle_delay_ms=0,
        action_wait_ms=1500,
    )

    assert [row["商品ID"] for row in rows] == ["100", "200", "300"]


def test_high_value_scan_waits_through_transient_disabled_next():
    class TransientLocator(FakePromotionLocator):
        def is_enabled(self):
            if self.selector == "#next" and self.page.disabled_reads:
                self.page.disabled_reads -= 1
                return False
            return super().is_enabled()

    page = FakePromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇，当前发布 1 篇"],
        ]
    )
    page.disabled_reads = 1
    page.locator = lambda selector: TransientLocator(page, selector)

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-07-30T10:00:00+08:00",
        filter_selector_key="high_value_filter",
        settle_delay_ms=0,
        action_wait_ms=1000,
    )

    assert [row["商品ID"] for row in rows] == ["100", "200"]


def test_high_value_scan_rejects_disabled_next_before_observed_terminal():
    class WrongTerminalLocator(FakePromotionLocator):
        def inner_text(self):
            if self.selector == "#terminal":
                return "2"
            return super().inner_text()

        def is_enabled(self):
            if self.selector == "#next":
                return False
            return super().is_enabled()

    page = FakePromotionPage(
        [["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"]]
    )
    page.locator = lambda selector: WrongTerminalLocator(page, selector)

    with pytest.raises(
        PaginationStateError, match="PAGINATION_TERMINAL_UNVERIFIED"
    ):
        scan_recommended_material_status(
            page,
            {
                "promotion_tab": "#promotion",
                "high_value_filter": "#recommended",
                "promotion_rows": ".promotion-row",
                "promotion_current_page": "#current",
                "promotion_first_page": "#first",
                "promotion_terminal_page": "#terminal",
                "promotion_next_page": "#next",
            },
            collected_at="2026-07-30T10:00:00+08:00",
            filter_selector_key="high_value_filter",
            settle_delay_ms=0,
            action_wait_ms=500,
        )


def test_high_value_scan_accepts_verified_single_page():
    page = FakePromotionPage(
        [["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"]]
    )
    events = []

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-07-30T10:00:00+08:00",
        filter_selector_key="high_value_filter",
        settle_delay_ms=0,
        action_wait_ms=500,
        on_pagination_event=events.append,
    )

    assert [row["商品ID"] for row in rows] == ["100"]
    assert [event["event_type"] for event in events] == [
        "origin",
        "terminal",
    ]


def test_high_value_scan_rejects_skipped_page_transition():
    class SkippingLocator(FakePromotionLocator):
        def click(self, **_kwargs):
            if self.selector == "#next":
                self.page.clicked.append(self.selector)
                self.page.page_index += 2
                return
            super().click(**_kwargs)

    page = FakePromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇，当前发布 1 篇"],
            ["商品三 商品ID 300 发布坑位到 9 篇，当前发布 2 篇"],
        ]
    )
    page.locator = lambda selector: SkippingLocator(page, selector)

    with pytest.raises(
        PaginationStateError, match="PAGINATION_TRANSITION_MISMATCH"
    ):
        scan_recommended_material_status(
            page,
            {
                "promotion_tab": "#promotion",
                "high_value_filter": "#recommended",
                "promotion_rows": ".promotion-row",
                "promotion_current_page": "#current",
                "promotion_first_page": "#first",
                "promotion_terminal_page": "#terminal",
                "promotion_next_page": "#next",
            },
            collected_at="2026-07-30T10:00:00+08:00",
            filter_selector_key="high_value_filter",
            settle_delay_ms=0,
            action_wait_ms=500,
        )


def test_promotion_scan_resumes_after_completed_pages_without_reemitting_them():
    page = FakePromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇、当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇、当前发布 1 篇"],
            ["商品三 商品ID 300 发布坑位到 9 篇、当前发布 2 篇"],
        ],
        page_index=2,
    )
    completed = [
        {
            "商品ID": "100",
            "目标容量": "9",
            "目标坑位": "9",
            "现有素材数": "0",
            "缺失数量": "9",
        },
        {
            "商品ID": "200",
            "目标容量": "9",
            "目标坑位": "9",
            "现有素材数": "1",
            "缺失数量": "8",
        },
    ]
    checkpoints = []

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-07-29T10:00:00+08:00",
        filter_selector_key="high_value_filter",
        initial_rows=completed,
        skip_completed_pages=2,
        expected_page_hashes={
            1: hashlib.sha256(b"100").hexdigest(),
            2: hashlib.sha256(b"200").hexdigest(),
        },
        on_page=lambda page_number, values: checkpoints.append(
            (page_number, [row["商品ID"] for row in values])
        ),
        settle_delay_ms=0,
    )

    assert page.clicked == [
        "#promotion",
        "#recommended",
        "#first",
        "#next",
        "#next",
    ]
    assert [row["商品ID"] for row in rows] == ["100", "200", "300"]
    assert checkpoints == [(3, ["100", "200", "300"])]


def test_export_reports_preserves_both_downloads_and_hashes(tmp_path):
    source_basic = tmp_path / "source-basic.xlsx"
    source_search = tmp_path / "source-search.xlsx"
    from openpyxl import Workbook

    basic_workbook = Workbook()
    basic_sheet = basic_workbook.active
    basic_sheet.append(["商品ID", "商品标题", "商品白底图", "短标题"])
    basic_sheet.append(["123", "儿童帽", "", ""])
    basic_workbook.save(source_basic)
    search_workbook = Workbook()
    search_sheet = search_workbook.active
    search_sheet.append(["商品ID", "素材类型", "素材ID", "审核状态"])
    search_sheet.append(["123", "图文", "M-1", "审核通过"])
    search_workbook.save(source_search)
    page = FakeExportPage(
        [
            FakeDownload(source_basic, "基础素材.xlsx"),
            FakeDownload(source_search, "推广素材.xlsx"),
        ]
    )

    records = export_reports(
        page,
        REQUIRED_SELECTOR_VALUES,
        tmp_path / "downloads",
        run_id="RUN-1",
        downloaded_at="2026-07-17T10:00:00+08:00",
        report_types=("basic", "promotion"),
    )

    assert len(records) == 2
    assert all(Path(record.path).exists() for record in records)
    assert all(len(record.sha256) == 64 for record in records)
    assert [record.report_type for record in records] == ["basic", "promotion"]
    assert [record.row_count for record in records] == [1, 1]
    assert all(record.schema_valid for record in records)
    assert records[0].contract_status == "confirmed"
    assert records[1].contract_status == "draft"
    assert Path(records[0].path).parent.name == "basic"
    assert Path(records[1].path).parent.name == "promotion"
    assert page.clicked == ["#export-basic", "#export-search"]


def test_export_reports_can_download_only_basic_materials(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "source-basic.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["商品ID", "商品标题", "商品白底图", "短标题"])
    sheet.append(["123", "儿童帽", "", ""])
    workbook.save(source)
    page = FakeExportPage([FakeDownload(source, "基础素材.xlsx")])

    records = export_reports(
        page,
        REQUIRED_SELECTOR_VALUES,
        tmp_path / "downloads",
        run_id="RUN-1",
        downloaded_at="2026-07-17T10:00:00+08:00",
        report_types=("basic",),
    )

    assert len(records) == 1
    assert records[0].report_type == "basic"
    assert records[0].row_count == 1
    assert records[0].schema_valid is True
    assert page.clicked == ["#export-basic"]


def test_export_reports_keeps_invalid_promotion_download_for_manual_review(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "source-promotion.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["商品ID", "未知字段"])
    sheet.append(["123", "value"])
    workbook.save(source)
    page = FakeExportPage([FakeDownload(source, "推广素材.xlsx")])

    records = export_reports(
        page,
        REQUIRED_SELECTOR_VALUES,
        tmp_path / "downloads",
        run_id="RUN-1",
        downloaded_at="2026-07-17T10:00:00+08:00",
        report_types=("promotion",),
    )

    assert Path(records[0].path).exists()
    assert records[0].schema_valid is False
    assert records[0].status == "needs_manual_review"
    assert records[0].reason_codes == ["PROMOTION_XLSX_SCHEMA_INVALID"]


def approved_item(tmp_path):
    image = tmp_path / "a.png"
    image.write_bytes(b"image")
    asset = AssetRecord(
        product_id="123",
        source_path=str(image),
        asset_type="image",
        license_status="confirmed",
        sha256=hashlib.sha256(b"image").hexdigest(),
        validation_status="valid",
    )
    return MaterialItem(
        task_id="MAT-1",
        product_id="123",
        material_type="image_text",
        slot_index=1,
        status=MaterialStatus.APPROVED,
        assets=[asset],
        title="KK树便携水杯",
        description="便携水杯设计，满足日常携带和饮水使用需求。",
        content_hash="fingerprint-1",
    )


def approval_manifest(item):
    return create_manifest(
        "KK Tree",
        [item],
        "operator",
        "2026-07-17T10:00:00+08:00",
        valid_until="2026-07-18T10:00:00+08:00",
    )


def configured_upload_page():
    page = FakePage()
    page.texts.update(
        {
            "#store": "KK Tree",
            "#product-id": "123",
            "#desired-slots": "3坑",
            "#material-table": "table",
        }
    )
    page.visible.add("#material-table")
    page.lists[".empty-slot"] = ["坑位1", "坑位2", "坑位3"]
    return page


def test_publish_timeout_does_not_click_twice(tmp_path):
    page = configured_upload_page()
    page.click_errors["#publish"] = PlaywrightTimeoutError("timeout")
    item = approved_item(tmp_path)

    outcome = upload_approved_item(
        page,
        item,
        approval_manifest(item),
        REQUIRED_SELECTOR_VALUES,
        expected_store="KK Tree",
        now="2026-07-17T11:00:00+08:00",
    )

    assert page.clicked.count("#publish") == 1
    assert outcome.status == "publish_uncertain"
    assert outcome.retry_allowed is False


def test_before_publish_checkpoint_runs_before_the_single_click(tmp_path):
    page = configured_upload_page()
    page.visible.add("#success")
    page.texts["#remote-id"] = "RM-1"
    item = approved_item(tmp_path)
    checkpoints = []

    def checkpoint():
        assert "#publish" not in page.clicked
        checkpoints.append("uploading-persisted")

    outcome = upload_approved_item(
        page,
        item,
        approval_manifest(item),
        REQUIRED_SELECTOR_VALUES,
        expected_store="KK Tree",
        now="2026-07-17T11:00:00+08:00",
        before_publish=checkpoint,
    )

    assert checkpoints == ["uploading-persisted"]
    assert page.clicked.count("#publish") == 1
    assert outcome.status == "submitted"


def test_changed_approved_content_stops_before_publish(tmp_path):
    page = configured_upload_page()
    item = approved_item(tmp_path)
    manifest = approval_manifest(item)
    item.title = "审批后修改"

    outcome = upload_approved_item(
        page,
        item,
        manifest,
        REQUIRED_SELECTOR_VALUES,
        expected_store="KK Tree",
        now="2026-07-17T11:00:00+08:00",
    )

    assert outcome.status == "blocked"
    assert outcome.reason == "APPROVED_CONTENT_CHANGED"
    assert "#publish" not in page.clicked


def test_replaced_asset_bytes_stop_before_publish(tmp_path):
    page = configured_upload_page()
    item = approved_item(tmp_path)
    manifest = approval_manifest(item)
    Path(item.assets[0].source_path).write_bytes(b"replaced after approval")

    outcome = upload_approved_item(
        page,
        item,
        manifest,
        REQUIRED_SELECTOR_VALUES,
        expected_store="KK Tree",
        now="2026-07-17T11:00:00+08:00",
    )

    assert outcome.status == "blocked"
    assert outcome.reason == "APPROVED_CONTENT_CHANGED"
    assert "#publish" not in page.clicked


def test_successful_publish_records_remote_id(tmp_path):
    page = configured_upload_page()
    page.visible.add("#success")
    page.texts["#remote-id"] = "RM-1"
    item = approved_item(tmp_path)

    outcome = upload_approved_item(
        page,
        item,
        approval_manifest(item),
        REQUIRED_SELECTOR_VALUES,
        expected_store="KK Tree",
        now="2026-07-17T11:00:00+08:00",
    )

    assert outcome.status == "submitted"
    assert outcome.remote_material_id == "RM-1"
    assert page.input_files == [("#file-input", [item.assets[0].source_path])]


def test_remote_match_resolves_uncertain(tmp_path):
    page = configured_upload_page()
    page.visible.add("#remote-table")
    page.texts.update(
        {
            "#remote-id": "RM-1",
            "#remote-fingerprint": "fingerprint-1",
            "#remote-status": "审核中",
            "#remote-slot": "1",
            "#remote-time": "2026-07-17T11:00:05+08:00",
        }
    )
    item = approved_item(tmp_path)

    outcome = verify_remote_item(
        page,
        item,
        REQUIRED_SELECTOR_VALUES,
        expected_store="KK Tree",
        submitted_at="2026-07-17T11:00:00+08:00",
    )

    assert outcome.status == "under_review"
    assert outcome.remote_material_id == "RM-1"
    assert outcome.retry_allowed is False
