from contextlib import contextmanager
import hashlib
import json
from pathlib import Path

import pytest

from upload_search_materials.browser.config import SelectorConfigError, load_selectors
from upload_search_materials.browser.export_page import export_reports
from upload_search_materials.browser.material_page import (
    PaginationStateError,
    SelectorInvalidError,
    _click_with_popup_retries,
    _settle_safe_popups,
    _wait_for_nonempty_promotion_rows,
    scan_recommended_material_status,
    supplement_material_status,
)
from upload_search_materials.browser import material_page as material_page_module
from upload_search_materials.browser.session import (
    HumanCheckRequired,
    StoreIdentityError,
    assert_store_identity,
    detect_human_check,
    prepare_collection_page,
    recommendation_material_center_url,
)
from upload_search_materials.browser.upload_page import upload_approved_item
from upload_search_materials.browser.verifier import verify_remote_item
from upload_search_materials.approval import create_manifest
from upload_search_materials.models import AssetRecord, MaterialItem, MaterialStatus


def test_false_success_fixture_matches_previous_final_page():
    fixture = json.loads(
        (
            Path(__file__).parent
            / "fixtures"
            / "session_20260730_032916_pagination_origin_sanitized.json"
        ).read_text(encoding="utf-8")
    )

    assert fixture["preceding_collection"]["row_count"] == 255
    assert fixture["preceding_collection"]["page_count"] == 26
    assert (
        fixture["preceding_collection"]["final_page_product_ids"]
        == fixture["false_success_collection"]["all_product_ids"]
    )


def test_prepare_collection_page_opens_recommend_tab_before_collection():
    class CollectionPage:
        def __init__(self):
            self.url = "https://myseller.taobao.com/home.htm"
            self.visited = []

        def goto(self, url, **kwargs):
            self.visited.append((url, kwargs))
            self.url = url

    page = CollectionPage()
    result = prepare_collection_page(
        page,
        "https://myseller.taobao.com/home.htm/material-center/material-management",
        {},
    )

    assert recommendation_material_center_url(page.url).endswith("?tab=recommend")
    assert page.visited[0][0].endswith("?tab=recommend")
    assert page.visited[0][1]["wait_until"] == "domcontentloaded"
    assert result["navigated"] is True


def test_prepare_collection_page_waits_for_spa_surface_before_popup_settlement(
    monkeypatch,
):
    events = []

    class Surface:
        first = None

        def __init__(self):
            self.first = self

        def wait_for(self, **kwargs):
            events.append(("surface", kwargs))

    class CollectionPage:
        url = "https://myseller.taobao.com/home.htm"

        def goto(self, url, **_kwargs):
            self.url = url

        def locator(self, _selector):
            return Surface()

    monkeypatch.setattr(
        material_page_module,
        "_settle_safe_popups",
        lambda *_args, **_kwargs: events.append(("popups", {})) or 0,
    )

    prepare_collection_page(
        CollectionPage(),
        "https://myseller.taobao.com/home.htm/material-center/material-management",
        {"promotion_tab": "#promotion-tab"},
    )

    assert events == [
        ("surface", {"state": "attached", "timeout": 15_000}),
        ("popups", {}),
    ]


def test_table_hydration_keeps_settling_late_popups(monkeypatch):
    settled = []

    class Rows:
        def __init__(self):
            self.reads = 0

        def all_inner_texts(self):
            self.reads += 1
            return [] if self.reads == 1 else ["商品ID 10001"]

    class Page:
        def __init__(self):
            self.rows = Rows()
            self.waited = []

        def locator(self, _selector):
            return self.rows

        def wait_for_timeout(self, milliseconds):
            self.waited.append(milliseconds)

    monkeypatch.setattr(
        material_page_module,
        "_settle_safe_popups",
        lambda _page, selectors, *, delay_ms: settled.append(
            (dict(selectors), delay_ms)
        ),
    )
    page = Page()

    rows = _wait_for_nonempty_promotion_rows(
        page,
        ".promotion-row",
        timeout_ms=1_000,
        popup_selectors={"safe_popup_close_priority": ".modal-close"},
    )

    assert rows == ["商品ID 10001"]
    assert len(settled) == 2
    assert page.waited == [250]
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError


REQUIRED_SELECTOR_VALUES = {
    "store_name": "#store",
    "human_check": "#human-check",
    "export_basic": "#export-basic",
    "export_search": "#export-search",
    "product_search": "#product-search",
    "product_id": "#product-id",
    "desired_slots": "#desired-slots",
    "material_table": "#material-table",
    "material_rows": ".material-row",
    "empty_slots": ".empty-slot",
    "review_status": ".review-status",
    "image_text_action": "#image-text-action",
    "video_action": "#video-action",
    "file_input": "#file-input",
    "title_input": "#title-input",
    "description_input": "#description-input",
    "publish_button": "#publish",
    "success_signal": "#success",
    "remote_material_id": "#remote-id",
    "remote_material_table": "#remote-table",
    "remote_material_fingerprint": "#remote-fingerprint",
    "remote_material_status": "#remote-status",
    "remote_material_slot": "#remote-slot",
    "remote_material_time": "#remote-time",
}


class FakeLocator:
    def __init__(self, page, selector):
        self.page = page
        self.selector = selector

    def inner_text(self):
        return self.page.texts.get(self.selector, "")

    def all_inner_texts(self):
        return list(self.page.lists.get(self.selector, []))

    def count(self):
        if self.selector in self.page.counts:
            return self.page.counts[self.selector]
        return 1 if self.selector in self.page.texts else 0

    def is_visible(self):
        return self.selector in self.page.visible

    def fill(self, value):
        self.page.filled.append((self.selector, value))

    def set_input_files(self, values):
        self.page.input_files.append((self.selector, list(values)))

    def press(self, key):
        self.page.pressed.append((self.selector, key))

    def click(self):
        self.page.clicked.append(self.selector)
        error = self.page.click_errors.get(self.selector)
        if error:
            raise error


class FakePage:
    def __init__(self):
        self.texts = {}
        self.lists = {}
        self.counts = {}
        self.visible = set()
        self.filled = []
        self.pressed = []
        self.clicked = []
        self.input_files = []
        self.click_errors = {}

    def locator(self, selector):
        return FakeLocator(self, selector)


class FakeDownload:
    def __init__(self, source, suggested_filename):
        self.source = Path(source)
        self.suggested_filename = suggested_filename

    def save_as(self, destination):
        Path(destination).write_bytes(self.source.read_bytes())


class FakeDownloadInfo:
    def __init__(self, download):
        self.value = download


class FakeExportPage(FakePage):
    def __init__(self, downloads):
        super().__init__()
        self.downloads = iter(downloads)

    @contextmanager
    def expect_download(self):
        yield FakeDownloadInfo(next(self.downloads))


def write_selectors(path, values):
    import yaml

    path.write_text(yaml.safe_dump(values, allow_unicode=True), encoding="utf-8")


def test_missing_required_selector_fails_configuration(tmp_path):
    path = tmp_path / "selectors.yaml"
    values = dict(REQUIRED_SELECTOR_VALUES)
    del values["publish_button"]
    write_selectors(path, values)

    with pytest.raises(SelectorConfigError, match="publish_button"):
        load_selectors(path)


def test_promotion_export_selector_accepts_new_name_and_normalizes_legacy_alias(
    tmp_path,
):
    path = tmp_path / "selectors.yaml"
    values = dict(REQUIRED_SELECTOR_VALUES)
    values["export_promotion"] = values.pop("export_search")
    write_selectors(path, values)

    selectors = load_selectors(path)

    assert selectors["export_promotion"] == "#export-search"
    assert selectors["export_search"] == "#export-search"


def test_missing_both_promotion_export_selector_names_fails_configuration(tmp_path):
    path = tmp_path / "selectors.yaml"
    values = dict(REQUIRED_SELECTOR_VALUES)
    values.pop("export_search")
    write_selectors(path, values)

    with pytest.raises(SelectorConfigError, match="export_promotion"):
        load_selectors(path)


def test_wrong_store_stops_batch():
    page = FakePage()
    page.texts["#store"] = "Other Store"

    with pytest.raises(StoreIdentityError):
        assert_store_identity(page, "#store", "KK Tree")


def test_store_identity_uses_exact_visible_match_among_wrappers():
    class Candidate:
        def __init__(self, text, visible=True):
            self.text = text
            self.visible = visible

        def inner_text(self):
            return self.text

        def is_visible(self):
            return self.visible

    class Candidates:
        def __init__(self, values):
            self.values = values

        def count(self):
            return len(self.values)

        def nth(self, index):
            return self.values[index]

    class StorePage:
        def locator(self, _selector):
            return Candidates(
                [
                    Candidate("kocotree旗舰店 5.0 88VIP"),
                    Candidate(""),
                    Candidate("kocotree旗舰店"),
                    Candidate("›"),
                ]
            )

    assert (
        assert_store_identity(
            StorePage(), "[class*=shopName]", "kocotree旗舰店"
        )
        == "kocotree旗舰店"
    )


def test_human_check_stops_batch():
    page = FakePage()
    page.visible.add("#human-check")

    with pytest.raises(HumanCheckRequired):
        detect_human_check(page, "#human-check")


def test_missing_material_table_is_not_interpreted_as_zero():
    page = FakePage()
    page.texts["#product-id"] = "123"
    page.texts["#desired-slots"] = "3坑"
    selectors = dict(REQUIRED_SELECTOR_VALUES)

    rows = supplement_material_status(
        page,
        selectors,
        ["123"],
        collected_at="2026-07-17T10:00:00+08:00",
    )

    assert rows[0]["状态"] == "needs_manual_review"
    assert rows[0]["原因码"] == "SELECTOR_INVALID"
    assert rows[0]["现有素材数"] == ""
    assert rows[0]["空坑位"] == ""
    assert "failed_field=material_table" in rows[0]["证据"]


def test_material_status_supplement_reads_exact_product():
    page = FakePage()
    page.texts.update({"#product-id": "123", "#desired-slots": "3坑", "#material-table": "table"})
    page.visible.add("#material-table")
    page.counts[".material-row"] = 2
    page.lists[".empty-slot"] = ["坑位3"]
    page.lists[".review-status"] = ["审核通过", "审核中"]

    rows = supplement_material_status(
        page,
        REQUIRED_SELECTOR_VALUES,
        ["123"],
        collected_at="2026-07-17T10:00:00+08:00",
    )

    assert rows == [
        {
            "商品ID": "123",
            "目标坑位": "3",
                "现有素材数": "2",
                "空坑位": "3",
                "精确坑位状态": "collected",
                "审核状态": "审核通过;审核中",
            "审核状态完整": "true",
            "状态": "ready_for_review",
            "原因码": "",
            "采集时间": "2026-07-17T10:00:00+08:00",
            "证据": "product=123;rows=2;selector_version=runtime",
        }
    ]


class FakePromotionLocator:
    def __init__(self, page, selector):
        self.page = page
        self.selector = selector

    def click(self, **_kwargs):
        self.page.clicked.append(self.selector)
        if self.selector == "#next":
            self.page.page_index += 1
        elif self.selector == "#first":
            self.page.page_index = 0
        elif self.selector == "#recommended":
            self.page.recommended_checked = True

    def count(self):
        if self.selector == ".promotion-row":
            return len(self.page.pages[self.page.page_index])
        return 1

    def all_inner_texts(self):
        if self.selector == ".promotion-row":
            return self.page.pages[self.page.page_index]
        return []

    def get_attribute(self, name):
        if self.selector == "#recommended" and name == "aria-checked":
            return "true" if self.page.recommended_checked else "false"
        return None

    def inner_text(self):
        if self.selector == "#recommended":
            return f"搜推高价值 {self.page.high_value_total}"
        if self.selector == "#current":
            return str(self.page.page_index + 1)
        if self.selector == "#first":
            return "1"
        if self.selector == "#terminal":
            return str(len(self.page.pages))
        return ""

    def is_enabled(self):
        return self.page.page_index < len(self.page.pages) - 1


class FakePromotionPage:
    def __init__(self, pages, *, page_index=0):
        self.pages = pages
        self.page_index = page_index
        self.high_value_total = sum(len(values) for values in pages)
        self.filtered_product_id = ""
        self.reload_count = 0
        self.recommended_checked = False
        self.clicked = []
        self.waited = []

    def locator(self, selector):
        return FakePromotionLocator(self, selector)

    def reload(self, **_kwargs):
        self.reload_count += 1
        self.filtered_product_id = ""

    def wait_for_timeout(self, milliseconds):
        self.waited.append(milliseconds)


def test_collection_skips_redundant_promotion_click_on_recommend_route():
    page = FakePromotionPage(
        [["商品一 商品ID 123 发布坑位到 9 篇，当前发布 0 篇"]]
    )
    page.url = (
        "https://myseller.taobao.com/home.htm/material-center/"
        "material-management?tab=recommend"
    )
    page.recommended_checked = True

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_next_page": "#next",
            "promotion_first_page": "#first",
            "promotion_current_page": "#current",
            "promotion_terminal_page": "#terminal",
        },
        collected_at="2026-08-03T00:00:00+00:00",
        filter_selector_key="high_value_filter",
        settle_delay_ms=0,
        action_wait_ms=0,
    )

    assert len(rows) == 1
    assert "#promotion" not in page.clicked


class DelayedPromotionLocator(FakePromotionLocator):
    def click(self, **_kwargs):
        if self.selector == "#next":
            self.page.clicked.append(self.selector)
            self.page.pending_next = True
            return
        super().click(**_kwargs)


class DelayedPromotionPage(FakePromotionPage):
    def __init__(self, pages):
        super().__init__(pages)
        self.pending_next = False

    def locator(self, selector):
        return DelayedPromotionLocator(self, selector)

    def wait_for_timeout(self, milliseconds):
        super().wait_for_timeout(milliseconds)
        if self.pending_next:
            self.page_index += 1
            self.pending_next = False


class SlowHydrationPromotionLocator(FakePromotionLocator):
    def count(self):
        if self.selector == ".promotion-row" and not self.page.rows_ready:
            return 0
        return super().count()

    def all_inner_texts(self):
        if self.selector == ".promotion-row" and not self.page.rows_ready:
            return []
        return super().all_inner_texts()


class SlowHydrationPromotionPage(FakePromotionPage):
    def __init__(self, pages):
        super().__init__(pages)
        self.rows_ready = False
        self.total_waited_ms = 0

    def locator(self, selector):
        return SlowHydrationPromotionLocator(self, selector)

    def wait_for_timeout(self, milliseconds):
        super().wait_for_timeout(milliseconds)
        self.total_waited_ms += milliseconds
        if self.total_waited_ms >= 4_000:
            self.rows_ready = True


class FakePopupLocator:
    def __init__(self, page, selector):
        self.page = page
        self.selector = selector

    def count(self):
        return 1 if self.page.overlay_open else 0

    def nth(self, _index):
        return self

    def is_visible(self):
        return self.page.overlay_open

    def inner_text(self):
        return (
            f"guide-step-{self.page.guide_steps}"
            if self.selector == "#guide-next"
            else "popup-close"
        )

    def click(self, **_kwargs):
        self.page.clicked.append(self.selector)
        if self.selector == "#guide-next":
            self.page.guide_steps -= 1
            if self.page.guide_steps == 0:
                self.page.overlay_open = False
            return
        if self.page.click_races and self.selector == "#opened-overlay-close":
            self.page.click_races -= 1
            raise PlaywrightTimeoutError("popup rerendered")
        if self.selector == "#opened-overlay-close":
            self.page.overlay_open = False


class FakePopupPage:
    def __init__(self, *, click_races=0, guide_steps=0):
        self.overlay_open = True
        self.click_races = click_races
        self.guide_steps = guide_steps
        self.clicked = []
        self.waited = []

    def locator(self, selector):
        return FakePopupLocator(self, selector)

    def wait_for_timeout(self, milliseconds):
        self.waited.append(milliseconds)


def test_popup_settle_prioritizes_open_overlay_close_control():
    page = FakePopupPage()

    closed = _settle_safe_popups(
        page,
        {
            "safe_popup_close_priority": "#opened-overlay-close",
            "safe_popup_close": "#generic-close",
        },
        delay_ms=0,
    )

    assert closed == 1
    assert page.clicked == ["#opened-overlay-close"]
    assert page._tmall_collection_events[-1]["changed"] is True


def test_popup_settle_retries_when_overlay_rerenders_during_click():
    page = FakePopupPage(click_races=1)

    closed = _settle_safe_popups(
        page,
        {
            "safe_popup_close_priority": "#opened-overlay-close",
        },
        delay_ms=0,
    )

    assert closed == 1
    assert page.clicked == [
        "#opened-overlay-close",
        "#opened-overlay-close",
    ]
    assert page._tmall_collection_events[-1]["click_mode"] == "force"


def test_popup_settle_watches_long_enough_for_a_late_first_install_guide():
    page = FakePopupPage(guide_steps=1)
    page.overlay_open = False
    original_wait = page.wait_for_timeout

    def delayed_open(milliseconds):
        original_wait(milliseconds)
        if len(page.waited) == 5:
            page.overlay_open = True

    page.wait_for_timeout = delayed_open

    closed = _settle_safe_popups(
        page,
        {"safe_popup_progress": "#guide-next"},
        delay_ms=10,
        quiet_checks_required=8,
    )

    assert closed == 1
    assert page.clicked == ["#guide-next"]


def test_popup_settle_closes_guide_before_advancing_steps():
    page = FakePopupPage(guide_steps=7)

    closed = _settle_safe_popups(
        page,
        {
            "safe_popup_progress": "#guide-next",
            "safe_popup_close_priority": "#opened-overlay-close",
        },
        delay_ms=0,
    )

    assert closed == 1
    assert page.clicked == ["#opened-overlay-close"]
    assert page.guide_steps == 7


def test_popup_settle_stops_one_noop_control_after_two_unchanged_actions():
    page = FakePopupPage()

    class NoOp(FakePopupLocator):
        def click(self, **_kwargs):
            self.page.clicked.append(self.selector)

    page.locator = lambda selector: NoOp(page, selector)

    closed = _settle_safe_popups(
        page,
        {"safe_popup_close_priority": "#noop"},
        delay_ms=0,
    )

    assert closed == 0
    assert page.clicked == ["#noop", "#noop"]


def test_popup_settle_catches_guide_that_appears_after_initial_quiet_check():
    page = FakePopupPage(guide_steps=1)
    page.overlay_open = False
    original_wait = page.wait_for_timeout

    def delayed_open(milliseconds):
        original_wait(milliseconds)
        if len(page.waited) == 1:
            page.overlay_open = True

    page.wait_for_timeout = delayed_open

    closed = _settle_safe_popups(
        page,
        {"safe_popup_progress": "#guide-next"},
        delay_ms=10,
    )

    assert closed == 1
    assert page.clicked == ["#guide-next"]


def test_popup_settle_closes_first_install_guide_inside_frame():
    page = FakePopupPage()
    page.overlay_open = False
    frame = FakePopupPage(guide_steps=1)
    page.frames = [frame]

    closed = _settle_safe_popups(
        page,
        {"safe_popup_progress": "#guide-next"},
        delay_ms=0,
    )

    assert closed == 1
    assert page.clicked == []
    assert frame.clicked == ["#guide-next"]
    assert page._tmall_collection_events[-1]["before"]["scope"] == "frame"


def test_popup_settle_continues_past_covered_control_in_stacked_dialogs():
    class StackedControl:
        def __init__(self, page, kind):
            self.page = page
            self.kind = kind

        def is_visible(self):
            return self.page.guide_open if self.kind == "guide" else self.page.ant_open

        def inner_text(self):
            return ""

        def get_attribute(self, name):
            return self.kind if name == "class" else ""

        def click(self, **_kwargs):
            self.page.clicked.append(self.kind)
            if self.kind == "ant":
                self.page.ant_open = False
            elif not self.page.ant_open:
                self.page.guide_open = False

    class StackedLocator:
        def __init__(self, page):
            self.page = page

        def values(self):
            output = []
            if self.page.guide_open:
                output.append(StackedControl(self.page, "guide"))
            if self.page.ant_open:
                output.append(StackedControl(self.page, "ant"))
            return output

        def count(self):
            return len(self.values())

        def nth(self, index):
            return self.values()[index]

    class StackedPage:
        def __init__(self):
            self.guide_open = True
            self.ant_open = True
            self.clicked = []
            self.frames = []

        def locator(self, _selector):
            return StackedLocator(self)

        def wait_for_timeout(self, _milliseconds):
            return None

    page = StackedPage()

    closed = _settle_safe_popups(
        page,
        {"safe_popup_close_priority": "#stacked-close"},
        delay_ms=0,
    )

    assert closed == 2
    assert page.ant_open is False
    assert page.guide_open is False
    assert page.clicked == ["guide", "ant", "guide"]


class OverlayBlockedTarget:
    def __init__(self):
        self.attempts = []
        self.evaluations = []

    def click(self, **kwargs):
        self.attempts.append(kwargs)
        if not kwargs.get("force"):
            raise PlaywrightTimeoutError("guide overlay intercepts pointer events")

    def evaluate(self, expression):
        self.evaluations.append(expression)


def test_successful_collection_click_does_not_scan_popup_controls():
    class DirectTarget:
        def __init__(self):
            self.attempts = []

        def click(self, **kwargs):
            self.attempts.append(kwargs)

    page = FakePopupPage()
    target = DirectTarget()

    _click_with_popup_retries(
        page,
        target,
        {
            "safe_popup_progress": "#guide-next",
            "safe_popup_close_priority": "#opened-overlay-close",
        },
        field_name="promotion_next_page",
        delay_ms=1000,
    )

    assert target.attempts == [{"timeout": 3000}]
    assert page.clicked == []
    assert page.waited == []


def test_read_only_collection_click_uses_dom_click_after_popup_retries():
    page = FakePopupPage()
    page.overlay_open = False
    target = OverlayBlockedTarget()

    _click_with_popup_retries(
        page,
        target,
        {"safe_popup_progress": "#guide-next"},
        field_name="high_value_filter",
        delay_ms=0,
    )

    assert len(target.attempts) == 2
    assert target.evaluations == ["element => element.click()"]
    assert page._tmall_collection_events == [
        {
            "action": "dom_click",
            "target_field": "high_value_filter",
            "reason": "recognized_guide_interceptor",
            "read_only": True,
        }
    ]


def test_unknown_overlay_never_uses_force_click():
    page = FakePopupPage()
    page.overlay_open = False
    target = OverlayBlockedTarget()

    with pytest.raises(SelectorInvalidError, match="popup_blocked"):
        _click_with_popup_retries(
            page,
            target,
            {},
            field_name="high_value_filter",
            delay_ms=0,
        )

    assert len(target.attempts) == 2


def test_write_target_never_uses_force_click_even_for_recognized_guide():
    page = FakePopupPage()
    page.overlay_open = False
    target = OverlayBlockedTarget()

    with pytest.raises(SelectorInvalidError, match="popup_blocked"):
        _click_with_popup_retries(
            page,
            target,
            {"safe_popup_progress": "#guide-next"},
            field_name="publish_button",
            delay_ms=0,
        )

    assert len(target.attempts) == 2


def test_recommended_promotion_scan_paginates_deduplicates_and_checkpoints():
    full_row = (
        "满坑商品 商品ID 565628742471 "
        "ID 3468934367960259 重复或图片有删除 "
        "ID 1539780903132488 ID 3295138149726183 ID 1787739833959372 "
        "ID 3212796032484488 ID 754653047484578 ID 506129500907 "
        "ID 461532917400 ID 461356707236 "
        "该商品为搜推高价值商品，已上调发布坑位到9篇，当前发布9篇。"
    )
    missing_row = (
        "缺口商品 商品ID 903588197784 "
        "ID 567934680336 ID 515805956420 ID 515848659407 "
        "该商品为搜推高价值商品，已上调发布坑位到9篇，当前发布3篇。"
    )
    unknown_target_row = "普通商品 商品ID 1037160921729 ID 557984331788"
    page = FakePromotionPage(
        [
            ["演示商品 商品ID xxxxxxxxxxxxxxx ID xxxxxxxxxxx1", full_row, missing_row],
            [missing_row, unknown_target_row],
        ]
    )
    checkpoints = []

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "recommended_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-07-24T01:57:30+08:00",
        on_page=lambda page_number, values: checkpoints.append(
            (page_number, [row["商品ID"] for row in values])
        ),
        settle_delay_ms=0,
    )

    assert page.clicked == ["#promotion", "#recommended", "#next"]
    assert [row["商品ID"] for row in rows] == [
        "565628742471",
        "903588197784",
        "1037160921729",
    ]
    assert rows[0]["目标容量"] == "9"
    assert rows[0]["现有素材数"] == "9"
    assert rows[0]["原因码"] == "REMOTE_MATERIAL_WARNING"
    assert rows[1]["缺失数量"] == "6"
    assert rows[1]["远端素材ID"] == "567934680336;515805956420;515848659407"
    assert rows[2]["目标容量"] == ""
    assert rows[2]["原因码"] == "TARGET_CAPACITY_NOT_EXPLICIT"
    assert checkpoints == [
        (1, ["565628742471", "903588197784"]),
        (2, ["565628742471", "903588197784", "1037160921729"]),
    ]


def test_promotion_scan_waits_until_next_page_product_ids_change():
    page = DelayedPromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇、当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇、当前发布 1 篇"],
        ]
    )

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-07-24T01:57:30+08:00",
        filter_selector_key="high_value_filter",
        settle_delay_ms=0,
        action_wait_ms=1000,
    )

    assert [row["商品ID"] for row in rows] == ["100", "200"]
    assert 250 in page.waited


def test_high_value_scan_allows_slow_next_page_hydration():
    class SlowTransitionPage(DelayedPromotionPage):
        def __init__(self, pages):
            super().__init__(pages)
            self.pending_waited_ms = 0

        def wait_for_timeout(self, milliseconds):
            self.waited.append(milliseconds)
            if not self.pending_next:
                return
            self.pending_waited_ms += milliseconds
            if self.pending_waited_ms >= 4_000:
                self.page_index += 1
                self.pending_next = False

    page = SlowTransitionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇、当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇、当前发布 1 篇"],
        ]
    )

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-07-24T01:57:30+08:00",
        filter_selector_key="high_value_filter",
        settle_delay_ms=0,
        action_wait_ms=500,
    )

    assert [row["商品ID"] for row in rows] == ["100", "200"]
    assert page.pending_waited_ms >= 4_000


@pytest.mark.parametrize("initial_page", [0, 1, 2])
def test_high_value_scan_normalizes_reused_tab_to_first_page(initial_page):
    page = FakePromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇，当前发布 1 篇"],
            ["商品三 商品ID 300 发布坑位到 9 篇，当前发布 2 篇"],
        ],
        page_index=initial_page,
    )
    events = []

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-07-30T10:00:00+08:00",
        filter_selector_key="high_value_filter",
        settle_delay_ms=0,
        action_wait_ms=1000,
        on_pagination_event=events.append,
    )

    assert [row["商品ID"] for row in rows] == ["100", "200", "300"]
    assert ("#first" in page.clicked) is (initial_page > 0)
    assert events[0]["event_type"] == "origin"
    assert events[0]["current_page"] == 1
    assert events[-1]["event_type"] == "terminal"
    assert events[-1]["current_page"] == 3


def test_high_value_scan_refreshes_stale_product_filter_and_checks_total():
    page = FakePromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到9篇，当前发布0篇"],
            ["商品二 商品ID 200 发布坑位到9篇，当前发布1篇"],
        ]
    )
    page.filtered_product_id = "100"
    events = []

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-08-03T10:00:00+08:00",
        filter_selector_key="high_value_filter",
        settle_delay_ms=0,
        action_wait_ms=500,
        on_pagination_event=events.append,
    )

    assert page.reload_count == 1
    assert page.filtered_product_id == ""
    assert len(rows) == 2
    origin = next(event for event in events if event["event_type"] == "origin")
    assert origin["expected_product_count"] == 2


def test_high_value_scan_waits_for_slow_spa_table_hydration():
    page = SlowHydrationPromotionPage(
        [["商品一 商品ID 100 发布坑位到9篇，当前发布0篇"]]
    )

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-08-07T10:00:00+08:00",
        filter_selector_key="high_value_filter",
        settle_delay_ms=0,
        action_wait_ms=500,
    )

    assert [row["商品ID"] for row in rows] == ["100"]
    assert page.total_waited_ms >= 4_000


def test_high_value_scan_fails_when_collected_count_differs_from_label():
    page = FakePromotionPage(
        [["商品一 商品ID 100 发布坑位到9篇，当前发布0篇"]]
    )
    page.high_value_total = 2

    with pytest.raises(PaginationStateError, match="HIGH_VALUE_TOTAL_MISMATCH"):
        scan_recommended_material_status(
            page,
            {
                "promotion_tab": "#promotion",
                "high_value_filter": "#recommended",
                "promotion_rows": ".promotion-row",
                "promotion_current_page": "#current",
                "promotion_first_page": "#first",
                "promotion_terminal_page": "#terminal",
                "promotion_next_page": "#next",
            },
            collected_at="2026-08-03T10:00:00+08:00",
            filter_selector_key="high_value_filter",
            settle_delay_ms=0,
            action_wait_ms=500,
        )


def test_high_value_scan_fails_when_first_page_reset_does_not_move():
    page = FakePromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇，当前发布 1 篇"],
        ],
        page_index=1,
    )
    original_locator = page.locator

    class NoOpFirst(FakePromotionLocator):
        def click(self, **_kwargs):
            if self.selector == "#first":
                self.page.clicked.append(self.selector)
                return
            super().click(**_kwargs)

    page.locator = lambda selector: NoOpFirst(page, selector)

    with pytest.raises(
        PaginationStateError, match="PAGINATION_ORIGIN_RESET_FAILED"
    ):
        scan_recommended_material_status(
            page,
            {
                "promotion_tab": "#promotion",
                "high_value_filter": "#recommended",
                "promotion_rows": ".promotion-row",
                "promotion_current_page": "#current",
                "promotion_first_page": "#first",
                "promotion_terminal_page": "#terminal",
                "promotion_next_page": "#next",
            },
            collected_at="2026-07-30T10:00:00+08:00",
            filter_selector_key="high_value_filter",
            settle_delay_ms=0,
            action_wait_ms=500,
        )

    page.locator = original_locator


def test_high_value_scan_waits_for_first_page_rows_to_replace_final_rows():
    class StaleRowsLocator(FakePromotionLocator):
        def click(self, **_kwargs):
            if self.selector == "#first":
                self.page.clicked.append(self.selector)
                self.page.stale_rows = list(
                    self.page.pages[self.page.page_index]
                )
                self.page.stale_reads = 2
                self.page.page_index = 0
                return
            super().click(**_kwargs)

        def all_inner_texts(self):
            if (
                self.selector == ".promotion-row"
                and self.page.stale_reads > 0
            ):
                return self.page.stale_rows
            return super().all_inner_texts()

    page = FakePromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇，当前发布 1 篇"],
            ["商品三 商品ID 300 发布坑位到 9 篇，当前发布 2 篇"],
        ],
        page_index=2,
    )
    page.stale_rows = []
    page.stale_reads = 0
    page.locator = lambda selector: StaleRowsLocator(page, selector)
    original_wait = page.wait_for_timeout

    def settle_rows(milliseconds):
        original_wait(milliseconds)
        if page.stale_reads > 0:
            page.stale_reads -= 1

    page.wait_for_timeout = settle_rows

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-07-30T10:00:00+08:00",
        filter_selector_key="high_value_filter",
        settle_delay_ms=0,
        action_wait_ms=1500,
    )

    assert [row["商品ID"] for row in rows] == ["100", "200", "300"]


def test_high_value_scan_waits_through_transient_disabled_next():
    class TransientLocator(FakePromotionLocator):
        def is_enabled(self):
            if self.selector == "#next" and self.page.disabled_reads:
                self.page.disabled_reads -= 1
                return False
            return super().is_enabled()

    page = FakePromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇，当前发布 1 篇"],
        ]
    )
    page.disabled_reads = 1
    page.locator = lambda selector: TransientLocator(page, selector)

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-07-30T10:00:00+08:00",
        filter_selector_key="high_value_filter",
        settle_delay_ms=0,
        action_wait_ms=1000,
    )

    assert [row["商品ID"] for row in rows] == ["100", "200"]


def test_high_value_scan_rejects_disabled_next_before_observed_terminal():
    class WrongTerminalLocator(FakePromotionLocator):
        def inner_text(self):
            if self.selector == "#terminal":
                return "2"
            return super().inner_text()

        def is_enabled(self):
            if self.selector == "#next":
                return False
            return super().is_enabled()

    page = FakePromotionPage(
        [["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"]]
    )
    page.locator = lambda selector: WrongTerminalLocator(page, selector)

    with pytest.raises(
        PaginationStateError, match="PAGINATION_TERMINAL_UNVERIFIED"
    ):
        scan_recommended_material_status(
            page,
            {
                "promotion_tab": "#promotion",
                "high_value_filter": "#recommended",
                "promotion_rows": ".promotion-row",
                "promotion_current_page": "#current",
                "promotion_first_page": "#first",
                "promotion_terminal_page": "#terminal",
                "promotion_next_page": "#next",
            },
            collected_at="2026-07-30T10:00:00+08:00",
            filter_selector_key="high_value_filter",
            settle_delay_ms=0,
            action_wait_ms=500,
        )


def test_high_value_scan_accepts_verified_single_page():
    page = FakePromotionPage(
        [["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"]]
    )
    events = []

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-07-30T10:00:00+08:00",
        filter_selector_key="high_value_filter",
        settle_delay_ms=0,
        action_wait_ms=500,
        on_pagination_event=events.append,
    )

    assert [row["商品ID"] for row in rows] == ["100"]
    assert [event["event_type"] for event in events] == [
        "origin",
        "terminal",
    ]


def test_high_value_scan_runs_random_action_every_one_or_two_new_pages():
    page = FakePromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇，当前发布 1 篇"],
            ["商品三 商品ID 300 发布坑位到 9 篇，当前发布 2 篇"],
        ]
    )
    intervals = iter((2, 1, 2))
    actions = []

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-08-10T10:00:00+08:00",
        filter_selector_key="high_value_filter",
        settle_delay_ms=0,
        action_wait_ms=500,
        on_page=lambda _page_number, _rows: None,
        random_action=lambda page_number: actions.append(page_number)
        or {"status": "completed", "page_state_restored": True},
        random_interval_picker=lambda: next(intervals),
    )

    assert [row["商品ID"] for row in rows] == ["100", "200", "300"]
    assert actions == [2, 3]
    assert [
        event["page_number"]
        for event in page._tmall_collection_events
        if event.get("action") == "random_collection_action"
    ] == [2, 3]


def test_high_value_scan_stops_if_random_action_does_not_restore_page():
    page = FakePromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇，当前发布 0 篇"],
        ]
    )

    with pytest.raises(
        PaginationStateError, match="RANDOM_ACTION_PAGE_RESTORE_FAILED"
    ):
        scan_recommended_material_status(
            page,
            {
                "promotion_tab": "#promotion",
                "high_value_filter": "#recommended",
                "promotion_rows": ".promotion-row",
                "promotion_current_page": "#current",
                "promotion_first_page": "#first",
                "promotion_terminal_page": "#terminal",
                "promotion_next_page": "#next",
            },
            collected_at="2026-08-10T10:00:00+08:00",
            filter_selector_key="high_value_filter",
            settle_delay_ms=0,
            action_wait_ms=500,
            on_page=lambda _page_number, _rows: None,
            random_action=lambda _page_number: {
                "status": "skipped",
                "reason_code": "RANDOM_ACTION_PAGE_RESTORE_FAILED",
                "page_state_restored": False,
                "detail": "url_matches=true;product_order_matches=false",
            },
            random_interval_picker=lambda: 1,
        )


def test_high_value_scan_checks_human_verification_around_checkpoint():
    page = FakePromotionPage(
        [["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"]]
    )
    checkpoints = []
    checks = []

    scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-08-10T10:00:00+08:00",
        filter_selector_key="high_value_filter",
        settle_delay_ms=0,
        action_wait_ms=500,
        on_page=lambda page_number, _rows: checkpoints.append(page_number),
        human_check_waiter=lambda page_number, location: checks.append(
            (page_number, location, tuple(checkpoints))
        ),
    )

    assert checks[0] == (1, "before_page", ())
    assert checks[1] == (1, "after_checkpoint", (1,))
    assert checks[2] == (1, "before_pagination", (1,))


def test_high_value_scan_rejects_skipped_page_transition():
    class SkippingLocator(FakePromotionLocator):
        def click(self, **_kwargs):
            if self.selector == "#next":
                self.page.clicked.append(self.selector)
                self.page.page_index += 2
                return
            super().click(**_kwargs)

    page = FakePromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇，当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇，当前发布 1 篇"],
            ["商品三 商品ID 300 发布坑位到 9 篇，当前发布 2 篇"],
        ]
    )
    page.locator = lambda selector: SkippingLocator(page, selector)

    with pytest.raises(
        PaginationStateError, match="PAGINATION_TRANSITION_MISMATCH"
    ):
        scan_recommended_material_status(
            page,
            {
                "promotion_tab": "#promotion",
                "high_value_filter": "#recommended",
                "promotion_rows": ".promotion-row",
                "promotion_current_page": "#current",
                "promotion_first_page": "#first",
                "promotion_terminal_page": "#terminal",
                "promotion_next_page": "#next",
            },
            collected_at="2026-07-30T10:00:00+08:00",
            filter_selector_key="high_value_filter",
            settle_delay_ms=0,
            action_wait_ms=500,
        )


def test_promotion_scan_resumes_after_completed_pages_without_reemitting_them():
    page = FakePromotionPage(
        [
            ["商品一 商品ID 100 发布坑位到 9 篇、当前发布 0 篇"],
            ["商品二 商品ID 200 发布坑位到 9 篇、当前发布 1 篇"],
            ["商品三 商品ID 300 发布坑位到 9 篇、当前发布 2 篇"],
        ],
        page_index=2,
    )
    completed = [
        {
            "商品ID": "100",
            "目标容量": "9",
            "目标坑位": "9",
            "现有素材数": "0",
            "缺失数量": "9",
        },
        {
            "商品ID": "200",
            "目标容量": "9",
            "目标坑位": "9",
            "现有素材数": "1",
            "缺失数量": "8",
        },
    ]
    checkpoints = []

    rows = scan_recommended_material_status(
        page,
        {
            "promotion_tab": "#promotion",
            "high_value_filter": "#recommended",
            "promotion_rows": ".promotion-row",
            "promotion_current_page": "#current",
            "promotion_first_page": "#first",
            "promotion_terminal_page": "#terminal",
            "promotion_next_page": "#next",
        },
        collected_at="2026-07-29T10:00:00+08:00",
        filter_selector_key="high_value_filter",
        initial_rows=completed,
        skip_completed_pages=2,
        expected_page_hashes={
            1: hashlib.sha256(b"100").hexdigest(),
            2: hashlib.sha256(b"200").hexdigest(),
        },
        on_page=lambda page_number, values: checkpoints.append(
            (page_number, [row["商品ID"] for row in values])
        ),
        settle_delay_ms=0,
    )

    assert page.clicked == [
        "#promotion",
        "#recommended",
        "#first",
        "#next",
        "#next",
    ]
    assert [row["商品ID"] for row in rows] == ["100", "200", "300"]
    assert checkpoints == [(3, ["100", "200", "300"])]


def test_export_reports_preserves_both_downloads_and_hashes(tmp_path):
    source_basic = tmp_path / "source-basic.xlsx"
    source_search = tmp_path / "source-search.xlsx"
    from openpyxl import Workbook

    basic_workbook = Workbook()
    basic_sheet = basic_workbook.active
    basic_sheet.append(["商品ID", "商品标题", "商品白底图", "短标题"])
    basic_sheet.append(["123", "儿童帽", "", ""])
    basic_workbook.save(source_basic)
    search_workbook = Workbook()
    search_sheet = search_workbook.active
    search_sheet.append(["商品ID", "素材类型", "素材ID", "审核状态"])
    search_sheet.append(["123", "图文", "M-1", "审核通过"])
    search_workbook.save(source_search)
    page = FakeExportPage(
        [
            FakeDownload(source_basic, "基础素材.xlsx"),
            FakeDownload(source_search, "推广素材.xlsx"),
        ]
    )

    records = export_reports(
        page,
        REQUIRED_SELECTOR_VALUES,
        tmp_path / "downloads",
        run_id="RUN-1",
        downloaded_at="2026-07-17T10:00:00+08:00",
        report_types=("basic", "promotion"),
    )

    assert len(records) == 2
    assert all(Path(record.path).exists() for record in records)
    assert all(len(record.sha256) == 64 for record in records)
    assert [record.report_type for record in records] == ["basic", "promotion"]
    assert [record.row_count for record in records] == [1, 1]
    assert all(record.schema_valid for record in records)
    assert records[0].contract_status == "confirmed"
    assert records[1].contract_status == "draft"
    assert Path(records[0].path).parent.name == "basic"
    assert Path(records[1].path).parent.name == "promotion"
    assert page.clicked == ["#export-basic", "#export-search"]


def test_export_reports_can_download_only_basic_materials(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "source-basic.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["商品ID", "商品标题", "商品白底图", "短标题"])
    sheet.append(["123", "儿童帽", "", ""])
    workbook.save(source)
    page = FakeExportPage([FakeDownload(source, "基础素材.xlsx")])

    records = export_reports(
        page,
        REQUIRED_SELECTOR_VALUES,
        tmp_path / "downloads",
        run_id="RUN-1",
        downloaded_at="2026-07-17T10:00:00+08:00",
        report_types=("basic",),
    )

    assert len(records) == 1
    assert records[0].report_type == "basic"
    assert records[0].row_count == 1
    assert records[0].schema_valid is True
    assert page.clicked == ["#export-basic"]


def test_export_reports_keeps_invalid_promotion_download_for_manual_review(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "source-promotion.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["商品ID", "未知字段"])
    sheet.append(["123", "value"])
    workbook.save(source)
    page = FakeExportPage([FakeDownload(source, "推广素材.xlsx")])

    records = export_reports(
        page,
        REQUIRED_SELECTOR_VALUES,
        tmp_path / "downloads",
        run_id="RUN-1",
        downloaded_at="2026-07-17T10:00:00+08:00",
        report_types=("promotion",),
    )

    assert Path(records[0].path).exists()
    assert records[0].schema_valid is False
    assert records[0].status == "needs_manual_review"
    assert records[0].reason_codes == ["PROMOTION_XLSX_SCHEMA_INVALID"]


def approved_item(tmp_path):
    image = tmp_path / "a.png"
    image.write_bytes(b"image")
    asset = AssetRecord(
        product_id="123",
        source_path=str(image),
        asset_type="image",
        license_status="confirmed",
        sha256=hashlib.sha256(b"image").hexdigest(),
        validation_status="valid",
    )
    return MaterialItem(
        task_id="MAT-1",
        product_id="123",
        material_type="image_text",
        slot_index=1,
        status=MaterialStatus.APPROVED,
        assets=[asset],
        title="KK树便携水杯",
        description="便携水杯设计，满足日常携带和饮水使用需求。",
        content_hash="fingerprint-1",
    )


def approval_manifest(item):
    return create_manifest(
        "KK Tree",
        [item],
        "operator",
        "2026-07-17T10:00:00+08:00",
        valid_until="2026-07-18T10:00:00+08:00",
    )


def configured_upload_page():
    page = FakePage()
    page.texts.update(
        {
            "#store": "KK Tree",
            "#product-id": "123",
            "#desired-slots": "3坑",
            "#material-table": "table",
        }
    )
    page.visible.add("#material-table")
    page.lists[".empty-slot"] = ["坑位1", "坑位2", "坑位3"]
    return page


def test_publish_timeout_does_not_click_twice(tmp_path):
    page = configured_upload_page()
    page.click_errors["#publish"] = PlaywrightTimeoutError("timeout")
    item = approved_item(tmp_path)

    outcome = upload_approved_item(
        page,
        item,
        approval_manifest(item),
        REQUIRED_SELECTOR_VALUES,
        expected_store="KK Tree",
        now="2026-07-17T11:00:00+08:00",
    )

    assert page.clicked.count("#publish") == 1
    assert outcome.status == "publish_uncertain"
    assert outcome.retry_allowed is False


def test_before_publish_checkpoint_runs_before_the_single_click(tmp_path):
    page = configured_upload_page()
    page.visible.add("#success")
    page.texts["#remote-id"] = "RM-1"
    item = approved_item(tmp_path)
    checkpoints = []

    def checkpoint():
        assert "#publish" not in page.clicked
        checkpoints.append("uploading-persisted")

    outcome = upload_approved_item(
        page,
        item,
        approval_manifest(item),
        REQUIRED_SELECTOR_VALUES,
        expected_store="KK Tree",
        now="2026-07-17T11:00:00+08:00",
        before_publish=checkpoint,
    )

    assert checkpoints == ["uploading-persisted"]
    assert page.clicked.count("#publish") == 1
    assert outcome.status == "submitted"


def test_changed_approved_content_stops_before_publish(tmp_path):
    page = configured_upload_page()
    item = approved_item(tmp_path)
    manifest = approval_manifest(item)
    item.title = "审批后修改"

    outcome = upload_approved_item(
        page,
        item,
        manifest,
        REQUIRED_SELECTOR_VALUES,
        expected_store="KK Tree",
        now="2026-07-17T11:00:00+08:00",
    )

    assert outcome.status == "blocked"
    assert outcome.reason == "APPROVED_CONTENT_CHANGED"
    assert "#publish" not in page.clicked


def test_replaced_asset_bytes_stop_before_publish(tmp_path):
    page = configured_upload_page()
    item = approved_item(tmp_path)
    manifest = approval_manifest(item)
    Path(item.assets[0].source_path).write_bytes(b"replaced after approval")

    outcome = upload_approved_item(
        page,
        item,
        manifest,
        REQUIRED_SELECTOR_VALUES,
        expected_store="KK Tree",
        now="2026-07-17T11:00:00+08:00",
    )

    assert outcome.status == "blocked"
    assert outcome.reason == "APPROVED_CONTENT_CHANGED"
    assert "#publish" not in page.clicked


def test_successful_publish_records_remote_id(tmp_path):
    page = configured_upload_page()
    page.visible.add("#success")
    page.texts["#remote-id"] = "RM-1"
    item = approved_item(tmp_path)

    outcome = upload_approved_item(
        page,
        item,
        approval_manifest(item),
        REQUIRED_SELECTOR_VALUES,
        expected_store="KK Tree",
        now="2026-07-17T11:00:00+08:00",
    )

    assert outcome.status == "submitted"
    assert outcome.remote_material_id == "RM-1"
    assert page.input_files == [("#file-input", [item.assets[0].source_path])]


def test_qianniu_publish_workflow_uses_migrated_page_flow(
    tmp_path, monkeypatch
):
    import upload_search_materials.browser.upload_page as upload_module
    from upload_search_materials.browser.qianniu_upload import (
        QianniuPublishObservation,
    )

    page = configured_upload_page()
    item = approved_item(tmp_path)
    calls = []

    def fake_prepare(page_arg, item_arg, *, material_center_url):
        calls.append(("prepare", item_arg.task_id, material_center_url))
        return {"OLD-1", "OLD-2"}

    def fake_publish(
        page_arg,
        item_arg,
        *,
        material_center_url,
        before_publish,
        before_remote_ids,
    ):
        assert before_remote_ids == {"OLD-1", "OLD-2"}
        calls.append(("button-ready", item_arg.task_id))
        before_publish()
        calls.append(("clicked", item_arg.task_id))
        return QianniuPublishObservation(
            "submitted",
            "",
            remote_material_id="RM-QIANNIU-1",
            evidence="exact-slot-verified",
        )

    monkeypatch.setattr(
        upload_module, "prepare_qianniu_upload", fake_prepare
    )
    monkeypatch.setattr(
        upload_module, "publish_qianniu_once", fake_publish
    )

    outcome = upload_approved_item(
        page,
        item,
        approval_manifest(item),
        REQUIRED_SELECTOR_VALUES,
        expected_store="KK Tree",
        now="2026-07-17T11:00:00+08:00",
        before_publish=lambda: calls.append(("checkpoint", "MAT-1")),
        workflow="qianniu_recommend",
        material_center_url="https://example.invalid/material-center",
    )

    assert outcome.status == "submitted"
    assert outcome.remote_material_id == "RM-QIANNIU-1"
    assert calls == [
        (
            "prepare",
            "MAT-1",
            "https://example.invalid/material-center",
        ),
        ("button-ready", "MAT-1"),
        ("checkpoint", "MAT-1"),
        ("clicked", "MAT-1"),
    ]


def test_qianniu_prepare_failure_stops_before_checkpoint(
    tmp_path, monkeypatch
):
    import upload_search_materials.browser.upload_page as upload_module
    from upload_search_materials.browser.qianniu_upload import (
        QianniuUploadError,
    )

    page = configured_upload_page()
    item = approved_item(tmp_path)
    checkpoints = []

    def fail_prepare(*args, **kwargs):
        raise QianniuUploadError(
            "QIANNIU_MATERIAL_IDENTITY_AMBIGUOUS"
        )

    monkeypatch.setattr(
        upload_module, "prepare_qianniu_upload", fail_prepare
    )

    outcome = upload_approved_item(
        page,
        item,
        approval_manifest(item),
        REQUIRED_SELECTOR_VALUES,
        expected_store="KK Tree",
        now="2026-07-17T11:00:00+08:00",
        before_publish=lambda: checkpoints.append("called"),
        workflow="qianniu_recommend",
    )

    assert outcome.status == "blocked"
    assert outcome.reason == "QIANNIU_MATERIAL_IDENTITY_AMBIGUOUS"
    assert outcome.retry_allowed is True
    assert checkpoints == []


def test_qianniu_remote_id_uses_set_delta_not_requested_position(
    tmp_path, monkeypatch
):
    import upload_search_materials.browser.qianniu_upload as module

    item = approved_item(tmp_path)
    item.slot_index = 5
    monkeypatch.setattr(module, "_open_recommend_list", lambda *args: None)
    monkeypatch.setattr(module, "_find_product_row", lambda *args: object())
    monkeypatch.setattr(
        module,
        "_remote_ids_from_row",
        lambda row: (
            {"OLD-1", "OLD-2", "NEW-1"},
            {
                "OLD-1": (4, "old"),
                "OLD-2": (5, "old"),
                "NEW-1": (1, "审核中"),
            },
        ),
    )

    observation = module._observe_new_remote_item(
        object(),
        item,
        before_remote_ids={"OLD-1", "OLD-2"},
        material_center_url="https://example.invalid/material-center",
    )

    assert observation.status == "under_review"
    assert observation.remote_material_id == "NEW-1"
    assert "requested_slot=5" in observation.evidence
    assert "current_position=1" in observation.evidence


def test_qianniu_remote_id_delta_fails_closed_when_baseline_changes(
    tmp_path, monkeypatch
):
    import upload_search_materials.browser.qianniu_upload as module

    item = approved_item(tmp_path)
    monkeypatch.setattr(module, "_open_recommend_list", lambda *args: None)
    monkeypatch.setattr(module, "_find_product_row", lambda *args: object())
    monkeypatch.setattr(
        module,
        "_remote_ids_from_row",
        lambda row: (
            {"OLD-2", "NEW-1"},
            {"OLD-2": (2, "old"), "NEW-1": (1, "new")},
        ),
    )

    observation = module._observe_new_remote_item(
        object(),
        item,
        before_remote_ids={"OLD-1", "OLD-2"},
        material_center_url="https://example.invalid/material-center",
    )

    assert observation.status == "publish_uncertain"
    assert observation.reason_code == "QIANNIU_REMOTE_ID_DELTA_AMBIGUOUS"


def test_remote_match_resolves_uncertain(tmp_path):
    page = configured_upload_page()
    page.visible.add("#remote-table")
    page.texts.update(
        {
            "#remote-id": "RM-1",
            "#remote-fingerprint": "fingerprint-1",
            "#remote-status": "审核中",
            "#remote-slot": "1",
            "#remote-time": "2026-07-17T11:00:05+08:00",
        }
    )
    item = approved_item(tmp_path)

    outcome = verify_remote_item(
        page,
        item,
        REQUIRED_SELECTOR_VALUES,
        expected_store="KK Tree",
        submitted_at="2026-07-17T11:00:00+08:00",
    )

    assert outcome.status == "under_review"
    assert outcome.remote_material_id == "RM-1"


def test_qianniu_resume_uses_exact_slot_verifier(
    tmp_path, monkeypatch
):
    import upload_search_materials.browser.verifier as verifier_module
    from upload_search_materials.browser.qianniu_upload import (
        QianniuPublishObservation,
    )

    page = configured_upload_page()
    item = approved_item(tmp_path)
    calls = []

    def fake_verify(page_arg, item_arg, *, material_center_url):
        calls.append((item_arg.task_id, material_center_url))
        return QianniuPublishObservation(
            "under_review",
            "",
            remote_material_id="RM-QIANNIU-2",
            evidence="product-and-slot-match",
        )

    monkeypatch.setattr(
        verifier_module, "verify_qianniu_remote_item", fake_verify
    )

    outcome = verify_remote_item(
        page,
        item,
        REQUIRED_SELECTOR_VALUES,
        expected_store="KK Tree",
        submitted_at="2026-07-17T11:00:00+08:00",
        workflow="qianniu_recommend",
        material_center_url="https://example.invalid/material-center",
    )

    assert outcome.status == "under_review"
    assert outcome.remote_material_id == "RM-QIANNIU-2"
    assert calls == [
        ("MAT-1", "https://example.invalid/material-center")
    ]
    assert outcome.retry_allowed is False
